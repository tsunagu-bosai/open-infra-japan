#!/usr/bin/env python3

import argparse
import csv
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
DATASETS = ROOT / "catalog" / "datasets.csv"
REFERENCE_MUNICIPALITIES = ROOT / "data" / "reference" / "municipalities.csv"
PUBLIC_TOILET_SCHEMA = ROOT / "schema" / "standard" / "public-toilet" / "schema.csv"


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def detect_csv_encoding(path):
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp932",
        "utf-16",
    ]

    for encoding in encodings:
        try:
            with path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as f:
                reader = csv.reader(f)
                header = next(reader)
                rows = sum(1 for _ in reader)

            return encoding, header, rows

        except (UnicodeDecodeError, StopIteration):
            continue

    return None, [], 0


def inspect_csv(path):
    encoding, header, rows = detect_csv_encoding(path)

    if encoding is None:
        return {
            "status": "ERROR",
            "format": "csv",
            "encoding": "unknown",
            "rows": "",
            "columns": "",
            "header": "",
        }

    return {
        "status": "OK",
        "format": "csv",
        "encoding": encoding,
        "rows": rows,
        "columns": len(header),
        "header": " | ".join(header),
        "header_fields": header,
    }


def inspect_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return {
            "status": f"ERROR: openpyxl import failed: {exc}",
            "format": "xlsx",
            "encoding": "-",
            "rows": "",
            "columns": "",
            "header": "",
        }

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        iterator = worksheet.iter_rows(values_only=True)

        header_row = next(iterator, None)

        if header_row is None:
            workbook.close()
            return {
                "status": "EMPTY",
                "format": "xlsx",
                "encoding": "-",
                "rows": 0,
                "columns": 0,
                "header": "",
            }

        header = [
            "" if value is None else str(value)
            for value in header_row
        ]

        rows = sum(1 for _ in iterator)

        workbook.close()

        return {
            "status": "OK",
            "format": "xlsx",
            "encoding": "-",
            "rows": rows,
            "columns": len(header),
            "header": " | ".join(header),
            "header_fields": header,
        }

    except Exception as exc:
        return {
            "status": f"ERROR: {exc}",
            "format": "xlsx",
            "encoding": "-",
            "rows": "",
            "columns": "",
            "header": "",
        }



def load_standard_header():
    if not PUBLIC_TOILET_SCHEMA.exists():
        return []

    rows = read_csv(PUBLIC_TOILET_SCHEMA)
    return [
        row["name"].strip()
        for row in rows
        if row.get("name") and row["name"].strip()
    ]


def normalize_header_name(value):
    value = (value or "").strip()
    value = value.replace("\n", "")
    value = value.replace("　", "")
    value = value.replace("＿", "_")
    value = value.replace("（", "(").replace("）", ")")
    value = value.replace("車いす", "車椅子")
    value = value.replace("多機能トイレ数", "バリアフリートイレ数")
    value = value.replace("画像ライセンス", "画像_ライセンス")
    value = value.replace("高度の域", "高度の値")
    return value


def classify_schema(result, standard_header):
    if result["status"] != "OK":
        return "invalid-file"

    header = result.get("header_fields") or []
    if not header:
        return "invalid-file"

    first = header[0].lstrip().lower()
    if first.startswith("<!doctype html") or first.startswith("<html"):
        return "invalid-file"

    # Ignore accidental empty columns at the end.
    trimmed = list(header)
    while trimmed and not str(trimmed[-1]).strip():
        trimmed.pop()

    if (
        standard_header
        and header == standard_header
        and len(header) == len(standard_header)
    ):
        return "exact"

    if standard_header and trimmed == standard_header:
        return "header-variation"

    normalized = [normalize_header_name(v) for v in trimmed]
    normalized_standard = [
        normalize_header_name(v) for v in standard_header
    ]

    if standard_header and normalized == normalized_standard:
        return "header-variation"

    # Older municipal-standard public-toilet datasets used fields such as
    # 都道府県コード又は市区町村コード / NO / 住所 / 方書 / 多機能トイレ数.
    legacy_markers = {
        "都道府県コード又は市区町村コード",
        "市区町村コード",
        "NO",
        "住所",
        "方書",
        "多機能トイレ数",
    }
    normalized_set = set(normalized)
    marker_hits = len(legacy_markers & normalized_set)

    standard_set = set(normalized_standard)
    overlap = len(normalized_set & standard_set)
    overlap_ratio = overlap / len(normalized_set) if normalized_set else 0

    if marker_hits >= 2 and overlap_ratio >= 0.55:
        return "legacy-standard"

    # A current-standard-derived file may omit a small number of columns.
    if standard_header and len(normalized) >= 35 and overlap_ratio >= 0.85:
        return "header-variation"

    return "non-standard"

def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--prefecture",
        help="都道府県コード（例: 13）。省略時は全国",
    )

    parser.add_argument(
        "--dataset",
        required=True,
        help="データ種別（例: public-toilet）",
    )

    args = parser.parse_args()

    standard_header = load_standard_header()
    datasets = read_csv(DATASETS)
    reference_municipalities = read_csv(REFERENCE_MUNICIPALITIES)
    reference_map = {
        row["municipality_code"].strip(): row
        for row in reference_municipalities
    }

    targets = [
        row for row in datasets
        if row["dataset_type"].strip() == args.dataset
        and (not args.prefecture or row["prefecture_code"].strip() == args.prefecture)
        and (row.get("publication_status") or "unknown").strip() == "published"
    ]

    results = []
    seen_directories = set()

    for row in targets:
        code = row["municipality_code"].strip()
        reference = reference_map.get(code) if code else None
        prefecture_code = row["prefecture_code"].strip()

        if reference:
            reference_prefecture_code = reference["prefecture_code"].strip()
            if reference_prefecture_code != prefecture_code:
                print(
                    f"MISSING: {row['municipality_name']} "
                    f"(prefecture code mismatch in reference)"
                )
                continue
            prefecture_name = reference["prefecture_name"].strip()
            scope_dir = f"{code}-{reference['municipality_name'].strip()}"
        elif not code:
            prefecture_name = row["prefecture_name"].strip()
            scope_dir = "prefecture"
        else:
            print(f"MISSING: {row['municipality_name']} (not in reference)")
            continue

        directory = (
            RAW_DIR
            / f"{prefecture_code}-{prefecture_name}"
            / scope_dir
            / args.dataset
        )

        # Multiple catalog rows may point at the same scope directory.
        directory_key = str(directory)
        if directory_key in seen_directories:
            continue
        seen_directories.add(directory_key)

        target_name = row["municipality_name"] or row["prefecture_name"]
        if not directory.exists():
            print(f"MISSING: {target_name} ({directory.relative_to(ROOT)})")
            continue

        files = [path for path in directory.iterdir() if path.is_file()]
        if not files:
            print(f"MISSING: {target_name} (no files)")
            continue

        for path in sorted(files):
            suffix = path.suffix.lower()
            if suffix == ".csv":
                result = inspect_csv(path)
            elif suffix == ".xlsx":
                result = inspect_xlsx(path)
            else:
                result = {
                    "status": "UNSUPPORTED", "format": suffix.lstrip("."),
                    "encoding": "", "rows": "", "columns": "", "header": "",
                    "header_fields": [],
                }
            result["municipality_code"] = code
            result["municipality_name"] = target_name
            result["file"] = path.name
            result["schema"] = classify_schema(result, standard_header)
            results.append(result)

    print()
    print(f"Files inspected: {len(results)}")
    print()

    for result in results:
        print(
            f"{result['municipality_code']} "
            f"{result['municipality_name']} "
            f"{result['file']}"
        )
        print(
            f"  status   : {result['status']}"
        )
        print(
            f"  format   : {result['format']}"
        )
        print(
            f"  encoding : {result['encoding']}"
        )
        print(
            f"  schema   : {result['schema']}"
        )
        print(
            f"  rows     : {result['rows']}"
        )
        print(
            f"  columns  : {result['columns']}"
        )
        print(
            f"  header   : {result['header']}"
        )
        print()



if __name__ == "__main__":
    main()
