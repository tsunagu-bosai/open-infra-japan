#!/usr/bin/env python3

import csv
import re
from pathlib import Path
from openpyxl import load_workbook

from schema import FIELDS

ROOT = Path(__file__).resolve().parents[3]

INPUT = (
    ROOT
    / "data/raw/22-静岡県/22100-静岡市/barrier-free"
    / "shizuoka-barrier-free.xlsx"
)

OUTPUT = (
    ROOT
    / "data/normalized/22-静岡県/22100-静岡市/barrier-free"
    / "shizuoka-barrier-free-facilities.csv"
)

SOURCE_NAME = "静岡市ユニバーサルデザイン・バリアフリーマップオープンデータ"
SOURCE_URL = "https://opendata.pref.shizuoka.jp/dataset/12491.html"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def header(value):
    return re.sub(r"\s+", "", clean(value).replace("\u3000", ""))


def presence(value, field_name, notes):
    value = clean(value)

    if not value:
        return ""

    if value == "有":
        return "yes"

    # 静岡データは「有＋詳細説明」の形が複数存在する。
    # 原文を失わずyesとして検索可能にする。
    notes.append(f"{field_name}={value}")
    return "yes"


WARD_OVERRIDES = {
    ("下川原南公園", "静岡市下川原南241"): "駿河区",
    ("大内観光トイレ", "清水大内599-1（霊山寺駐車場内）"): "清水区",
}


def ward_from_address(address, facility_name=""):
    address = clean(address)
    facility_name = clean(facility_name)

    for ward in ("葵区", "駿河区", "清水区"):
        if ward in address:
            return ward

    return WARD_OVERRIDES.get(
        (facility_name, address),
        "",
    )


def main():
    wb = load_workbook(INPUT, read_only=True, data_only=True)

    output_rows = []

    configs = [
        (
            "一覧表（バリアフリーマップ）",
            "barrier_free_map",
        ),
        (
            "一覧表（気軽にトイレマップ）",
            "casual_toilet_map",
        ),
    ]

    for sheet_name, dataset_name in configs:
        ws = wb[sheet_name]

        rows = ws.iter_rows(values_only=True)
        raw_headers = next(rows)
        headers = [header(v) for v in raw_headers]

        for source_row, row in enumerate(rows, start=2):
            values = {
                headers[i]: row[i]
                for i in range(len(headers))
                if headers[i]
            }

            facility_name = clean(values.get("施設名"))
            formal_name = clean(values.get("正式名称"))
            address = clean(values.get("住所"))

            # XLSX末尾に大量の空行が存在するため除外。
            if not facility_name and not formal_name and not address:
                continue

            notes = []

            wheelchair_toilet = presence(
                values.get("車いす対応トイレ"),
                "wheelchair_toilet",
                notes,
            )
            ostomate = presence(
                values.get("オストメイト対応設備"),
                "ostomate",
                notes,
            )
            adult_bed = presence(
                values.get("成人用シート"),
                "adult_bed",
                notes,
            )
            diaper = presence(
                values.get("おむつ替えシート"),
                "diaper_changing_space",
                notes,
            )
            child_toilet = presence(
                values.get("子ども用トイレ"),
                "child_toilet",
                notes,
            )
            child_chair = presence(
                values.get("子ども用チェアー"),
                "child_chair",
                notes,
            )

            nursing = presence(
                values.get("授乳室"),
                "nursing_space",
                notes,
            )
            stroller = presence(
                values.get("ベビーカー貸し出し"),
                "stroller_rental",
                notes,
            )

            out = {
                "prefecture_code": "22",
                "prefecture_name": "静岡県",
                "municipality_code": "22100",
                "municipality_name": "静岡市",
                "ward_name": ward_from_address(
                    address,
                    facility_name or formal_name,
                ),

                "facility_name": facility_name or formal_name,
                "formal_name": formal_name,
                "facility_category": clean(values.get("施設分類")),
                "postal_code": clean(values.get("郵便番号")),
                "address": address,
                "latitude": clean(values.get("緯度")),
                "longitude": clean(values.get("経度")),
                "homepage_url": clean(values.get("ホームページ")),

                "multipurpose_toilet": "",
                "wheelchair_toilet": wheelchair_toilet,
                "toilet_floor": "",
                "unisex_toilet_count": "",
                "male_toilet_count": "",
                "female_toilet_count": "",

                "washlet": "",
                "adult_bed": adult_bed,
                "ostomate": ostomate,
                "voice_guidance": "",
                "baby_bed": "",
                "baby_chair": "",
                "child_toilet": child_toilet,
                "child_chair": child_chair,
                "flush_type": "",
                "emergency_call_type": "",

                "male_urinal_with_rail": presence(
                    values.get("小便器手すり"),
                    "male_urinal_with_rail",
                    notes,
                ),
                "male_western_toilet": "",
                "female_western_toilet": "",
                "male_baby_bed": "",
                "female_baby_bed": "",
                "male_baby_chair": "",
                "female_baby_chair": "",

                # 「スロープ・段差なし」は entrance_step と意味が逆なので、
                # 誤変換せず今回は共通列には入れない。
                "entrance_step": "",
                "stroller_rental": stroller,
                "nursing_space": nursing,
                "diaper_changing_space": diaper,
                "kids_room": "",

                "source_dataset": dataset_name,
                "source_name": SOURCE_NAME,
                "source_url": SOURCE_URL,
                "source_updated_at": "",
                "source_row": str(source_row),
                "attribute_note": " | ".join(notes),
                "note": clean(values.get("備考（防災対応、駐車場台数）")),
            }

            output_rows.append(out)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=FIELDS,
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(output_rows)

    print("output rows:", len(output_rows))
    print("output:", OUTPUT)


if __name__ == "__main__":
    main()
