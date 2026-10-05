#!/usr/bin/env python3

import csv
from pathlib import Path
from openpyxl import load_workbook

from schema import FIELDS

ROOT = Path(__file__).resolve().parents[3]

INPUT = ROOT / "data/raw/26-京都府/barrier-free/kyoto-barrier-free-facilities.xlsx"
OUTPUT = ROOT / "data/normalized/26-京都府/barrier-free/kyoto-barrier-free-facilities.csv"

SOURCE_NAME = "京都府内のバリアフリー施設一覧"
SOURCE_URL = "https://www.pref.kyoto.jp/f-machi/hitoyasa.html"
SOURCE_UPDATED_AT = "2026-09"

MUNICIPALITY_CODES = {
    "京都市": "26100",
    "福知山市": "26201",
    "舞鶴市": "26202",
    "綾部市": "26203",
    "宇治市": "26204",
    "宮津市": "26205",
    "亀岡市": "26206",
    "城陽市": "26207",
    "向日市": "26208",
    "長岡京市": "26209",
    "八幡市": "26210",
    "京田辺市": "26211",
    "京丹後市": "26212",
    "南丹市": "26213",
    "木津川市": "26214",
    "大山崎町": "26303",
    "久御山町": "26322",
    "井手町": "26343",
    "宇治田原町": "26344",
    "笠置町": "26364",
    "和束町": "26365",
    "精華町": "26366",
    "南山城村": "26367",
    "京丹波町": "26407",
    "伊根町": "26463",
    "与謝野町": "26465",
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def state(value):
    v = clean(value)
    return {
        "あり": "yes",
        "なし": "no",
        "調査中": "unknown",
        "": "",
    }.get(v, v)


def flush_type(value):
    v = clean(value)
    return {
        "ボタン": "button",
        "レバー": "lever",
        "センサー": "sensor",
        "その他": "other",
        "なし": "none",
        "調査中": "unknown",
        "": "",
    }.get(v, v)


def emergency_call(value):
    v = clean(value)
    return {
        "ボタン式あり": "button",
        "引っ張り式あり": "pull_cord",
        "なし": "none",
        "調査中": "unknown",
        "": "",
    }.get(v, v)


def municipality_parts(raw):
    raw = clean(raw)

    if raw.startswith("京都市 "):
        return "京都市", raw.removeprefix("京都市 ").strip()

    return raw, ""


def main():
    wb = load_workbook(INPUT, read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    index = {name: i for i, name in enumerate(headers)}

    output_rows = []

    for source_row, row in enumerate(rows, start=2):
        municipality_name, ward_name = municipality_parts(
            row[index["市区町村 "]]
        )

        municipality_code = MUNICIPALITY_CODES.get(municipality_name)
        if not municipality_code:
            raise ValueError(
                f"unknown municipality: {municipality_name!r}"
            )

        out = {
            "prefecture_code": "26",
            "prefecture_name": "京都府",
            "municipality_code": municipality_code,
            "municipality_name": municipality_name,
            "ward_name": ward_name,

            "facility_name": clean(row[index["施設名"]]),
            "formal_name": "",
            "facility_category": "",
            "postal_code": clean(row[index["郵便番号"]]),
            "address": clean(row[index["所在地"]]),
            "latitude": "",
            "longitude": "",
            "homepage_url": clean(row[index["ホームページ (基本情報)"]]),

            "multipurpose_toilet": state(
                row[index["多目的トイレ (基本情報)"]]
            ),
            "wheelchair_toilet": state(
                row[index["車いす使用者用トイレ (多目的トイレ)"]]
            ),
            "toilet_floor": clean(
                row[index["設置階 (多目的トイレ)"]]
            ),
            "unisex_toilet_count": clean(
                row[index["個数（男女共用） (多目的トイレ)"]]
            ),
            "male_toilet_count": clean(
                row[index["個数（男子） (多目的トイレ)"]]
            ),
            "female_toilet_count": clean(
                row[index["個数（女子） (多目的トイレ)"]]
            ),

            "washlet": state(
                row[index["温水洗浄便座 (多目的トイレ)"]]
            ),
            "adult_bed": state(
                row[index["多目的ベッド (多目的トイレ)"]]
            ),
            "ostomate": state(
                row[index["オストメイト設備 (多目的トイレ)"]]
            ),
            "voice_guidance": state(
                row[index["音声案内装置 (多目的トイレ)"]]
            ),
            "baby_bed": state(
                row[index["ベビーベッド (多目的トイレ)"]]
            ),
            "baby_chair": state(
                row[index["ベビーチェア (多目的トイレ)"]]
            ),
            "child_toilet": "",
            "child_chair": "",
            "flush_type": flush_type(
                row[index["便器洗浄 (多目的トイレ)"]]
            ),
            "emergency_call_type": emergency_call(
                row[index["非常呼び出し装置 (多目的トイレ)"]]
            ),

            "male_urinal_with_rail": state(
                row[index["手すりつき小便器 (一般トイレ（男性用）)"]]
            ),
            "male_western_toilet": state(
                row[index["洋式便座 (一般トイレ（男性用）)"]]
            ),
            "female_western_toilet": state(
                row[index["洋式便座 (一般トイレ（女性用）)"]]
            ),
            "male_baby_bed": state(
                row[index["ベビーベッド (一般トイレ（男性用）)"]]
            ),
            "female_baby_bed": state(
                row[index["ベビーベッド (一般トイレ（女性用）)"]]
            ),
            "male_baby_chair": state(
                row[index["ベビーチェア (一般トイレ（男性用）)"]]
            ),
            "female_baby_chair": state(
                row[index["ベビーチェア (一般トイレ（女性用）)"]]
            ),

            "entrance_step": state(
                row[index["出入口に段差 (建物主要出入口)"]]
            ),
            "stroller_rental": state(
                row[index["ベビーカー貸出 (子育て支援)"]]
            ),
            "nursing_space": state(
                row[index["授乳スペース (子育て支援)"]]
            ),
            "diaper_changing_space": state(
                row[index["おむつ交換スペース (子育て支援)"]]
            ),
            "kids_room": state(
                row[index["キッズルーム (子育て支援)"]]
            ),

            "source_dataset": "barrier_free_facility_list",
            "source_name": SOURCE_NAME,
            "source_url": SOURCE_URL,
            "source_updated_at": SOURCE_UPDATED_AT,
            "source_row": str(source_row),
            "attribute_note": "",
            "note": clean(row[index["備考 (備考)"]]),
        }

        output_rows.append(out)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=FIELDS,
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(output_rows)

    print("input rows :", len(output_rows))
    print("output rows:", len(output_rows))
    print("output     :", OUTPUT)


if __name__ == "__main__":
    main()
