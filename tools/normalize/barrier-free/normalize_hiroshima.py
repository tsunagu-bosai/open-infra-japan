#!/usr/bin/env python3

import csv
from pathlib import Path
from openpyxl import load_workbook

from schema import FIELDS

ROOT = Path(__file__).resolve().parents[3]

INPUT = (
    ROOT
    / "data/raw/34-広島県/barrier-free"
    / "hiroshima-ostomate.xlsx"
)

OUTPUT = (
    ROOT
    / "data/normalized/34-広島県/barrier-free"
    / "hiroshima-ostomate.csv"
)

SOURCE_NAME = "広島県内のオストメイト対応トイレ一覧"
SOURCE_URL = "https://www.pref.hiroshima.lg.jp/soshiki/62/1246620261731.html"
SOURCE_UPDATED_AT = "2025-07-01"

MUNICIPALITY_CODES = {
    "広島市": "34100",
    "呉市": "34202",
    "竹原市": "34203",
    "三原市": "34204",
    "尾道市": "34205",
    "福山市": "34207",
    "府中市": "34208",
    "三次市": "34209",
    "庄原市": "34210",
    "大竹市": "34211",
    "東広島市": "34212",
    "廿日市市": "34213",
    "安芸高田市": "34214",
    "江田島市": "34215",
    "府中町": "34302",
    "海田町": "34304",
    "熊野町": "34307",
    "坂町": "34309",
    "安芸太田町": "34368",
    "北広島町": "34369",
    "大崎上島町": "34431",
    "世羅町": "34462",
    "神石高原町": "34545",
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def has_mark(value):
    value = clean(value)
    if not value:
        return ""
    return "yes"


def main():
    wb = load_workbook(INPUT, read_only=True, data_only=True)

    output_rows = []

    for ws in wb.worksheets:
        if ws.title == "表の説明":
            continue

        for source_row, row in enumerate(
            ws.iter_rows(min_row=5, values_only=True),
            start=5,
        ):
            municipality_name = clean(row[0])
            facility_name = clean(row[1])
            address = clean(row[4])

            if not municipality_name and not facility_name and not address:
                continue

            municipality_code = MUNICIPALITY_CODES.get(municipality_name)
            if not municipality_code:
                raise ValueError(
                    f"unknown municipality at {ws.title}:{source_row}: "
                    f"{municipality_name!r}"
                )

            attributes = []

            for key, value in [
                ("installation_count", row[3]),
                ("waste_sink", row[5]),
                ("warm_water_shower", row[6]),
                ("simple_faucet", row[7]),
                ("mirror", row[8]),
                ("changing_table_or_bed", row[9]),
                ("wheelchair_accessible", row[10]),
                ("ostomate_mark", row[11]),
            ]:
                v = clean(value)
                if v:
                    attributes.append(f"{key}={v}")

            out = {
                "prefecture_code": "34",
                "prefecture_name": "広島県",
                "municipality_code": municipality_code,
                "municipality_name": municipality_name,
                "ward_name": "",

                "facility_name": facility_name,
                "formal_name": "",
                "facility_category": clean(row[2]),
                "postal_code": "",
                "address": address,
                "latitude": "",
                "longitude": "",
                "homepage_url": "",

                "multipurpose_toilet": "",
                "wheelchair_toilet": has_mark(row[10]),
                "toilet_floor": "",
                "unisex_toilet_count": "",
                "male_toilet_count": "",
                "female_toilet_count": "",

                "washlet": "",
                "adult_bed": "",
                "ostomate": "yes",
                "voice_guidance": "",
                "baby_bed": "",
                "baby_chair": "",
                "child_toilet": "",
                "child_chair": "",
                "flush_type": "",
                "emergency_call_type": "",

                "male_urinal_with_rail": "",
                "male_western_toilet": "",
                "female_western_toilet": "",
                "male_baby_bed": "",
                "female_baby_bed": "",
                "male_baby_chair": "",
                "female_baby_chair": "",

                "entrance_step": "",
                "stroller_rental": "",
                "nursing_space": "",
                "diaper_changing_space": "",
                "kids_room": "",

                "source_dataset": "ostomate_list",
                "source_name": SOURCE_NAME,
                "source_url": SOURCE_URL,
                "source_updated_at": SOURCE_UPDATED_AT,
                "source_row": f"{ws.title}:{source_row}",
                "attribute_note": " | ".join(attributes),
                "note": clean(row[12]),
            }

            output_rows.append(out)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=FIELDS,
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(output_rows)

    print("input rows :", len(output_rows))
    print("output rows:", len(output_rows))
    print("output     :", OUTPUT)


if __name__ == "__main__":
    main()
