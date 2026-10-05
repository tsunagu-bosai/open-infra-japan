from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def read_xlsx_values(
    path: Path,
    sheet_name: str | None = None,
) -> list[tuple[Any, ...]]:
    """Read one XLSX worksheet as raw row values.

    This function owns only workbook/sheet I/O. Value normalization and
    source-specific interpretation remain the caller's responsibility.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"{path}: sheet not found: {sheet_name!r}")

        worksheet = (
            workbook[sheet_name]
            if sheet_name is not None
            else workbook[workbook.sheetnames[0]]
        )
        return list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
