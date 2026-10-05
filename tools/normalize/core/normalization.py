from __future__ import annotations

import csv
import hashlib
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from openpyxl import load_workbook

from tools.normalize.core.encoding import detect_text_encoding


@dataclass(frozen=True)
class SourceTable:
    header: list[str]
    rows: list[dict[str, str]]
    source_format: str


class RowSupport:
    @staticmethod
    def clean(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def clean_cell(value: Any) -> str:
        """Clean spreadsheet cells while preserving integer-looking IDs/counts."""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def has_meaningful_value(cls, values: Iterable[Any]) -> bool:
        return any(cls.clean(value) for value in values)

    @classmethod
    def append_text(
        cls,
        row: dict[str, str],
        field: str,
        text: str,
        *,
        separator: str = " / ",
    ) -> None:
        text = cls.clean(text)
        if not text:
            return
        current = cls.clean(row.get(field))
        if text in current:
            return
        row[field] = f"{current}{separator}{text}" if current else text

    @classmethod
    def append_note(
        cls,
        row: dict[str, str],
        text: str,
        *,
        separator: str = " / ",
    ) -> None:
        cls.append_text(row, "備考", text, separator=separator)

    @classmethod
    def append_labeled_note(
        cls,
        row: dict[str, str],
        label: str,
        value: Any,
    ) -> None:
        text = cls.clean(value)
        if text:
            cls.append_note(row, f"{label}={text}")


class ValueCleaner:
    """Reusable value-cleaning strategy with configurable placeholder values."""

    def __init__(self, placeholders: Iterable[str] = ()) -> None:
        self.placeholders = frozenset(str(value) for value in placeholders)

    def __call__(self, value: Any, *, preserve_placeholder: bool = False) -> str:
        text = RowSupport.clean(value)
        if not preserve_placeholder and text in self.placeholders:
            return ""
        return text


class AllowedValueStrategy:
    """Validate a cleaned scalar value against an explicit allowed-value set."""

    def __init__(
        self,
        allowed_values: Iterable[str],
        *,
        cleaner: Callable[[Any], str] = RowSupport.clean,
    ) -> None:
        self.allowed_values = frozenset(str(value) for value in allowed_values)
        self.cleaner = cleaner

    def __call__(self, value: Any, field: str) -> str:
        text = self.cleaner(value)
        if text not in self.allowed_values:
            raise RuntimeError(f"{field}: unexpected value {text!r}")
        return text


class TimeNormalizationStrategy(ABC):
    @abstractmethod
    def normalize(self, value: Any, *args: Any, **kwargs: Any) -> Any:
        raise NotImplementedError

    def __call__(self, value: Any, *args: Any, **kwargs: Any) -> Any:
        return self.normalize(value, *args, **kwargs)


class StrictHmsTimeStrategy(TimeNormalizationStrategy):
    """Normalize HH:MM or HH:MM:00 and reject malformed values."""

    def __init__(self, cleaner: Callable[[Any], str] = RowSupport.clean) -> None:
        self.cleaner = cleaner

    def normalize(self, value: Any) -> str:
        text = self.cleaner(value).replace("：", ":")
        if not text:
            return ""
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
        if not match:
            raise RuntimeError(f"unexpected time: {text!r}")
        hour = int(match.group(1))
        minute = int(match.group(2))
        second = match.group(3)
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            raise RuntimeError(f"invalid time: {text!r}")
        if second not in (None, "00"):
            raise RuntimeError(f"seconds are not zero: {text!r}")
        return f"{hour:02d}:{minute:02d}"


class PreserveInvalidTimeStrategy(TimeNormalizationStrategy):
    """Normalize valid times while returning the original invalid value for notes."""

    def normalize(self, value: Any) -> tuple[str, str | None]:
        text = RowSupport.clean(value).replace("：", ":")
        if not text:
            return "", None
        if text in {"24:00", "24:00:00"}:
            return "", text
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
        if not match:
            return "", text
        hour = int(match.group(1))
        minute = int(match.group(2))
        second = match.group(3)
        if hour > 23 or minute > 59 or second not in (None, "00"):
            return "", text
        return f"{hour:02d}:{minute:02d}", None


class LenientEndOfDayTimeStrategy(TimeNormalizationStrategy):
    """Normalize valid HH:MM, map 24:00 to 23:59, otherwise preserve text."""

    def __init__(self, cleaner: Callable[[Any], str] = RowSupport.clean) -> None:
        self.cleaner = cleaner

    def normalize(self, value: Any) -> str:
        text = self.cleaner(value)
        if not text:
            return ""
        if text == "24:00":
            return "23:59"
        parts = text.split(":")
        if len(parts) == 2 and all(part.isdigit() for part in parts):
            hour, minute = map(int, parts)
            if 0 <= hour <= 23 and 0 <= minute <= 59:
                return f"{hour:02d}:{minute:02d}"
        return text


class LenientHmsTimeStrategy(TimeNormalizationStrategy):
    """Normalize HH:MM / HH:MM:00 and preserve unsupported text unchanged."""

    def __init__(
        self,
        cleaner: Callable[[Any], str] = RowSupport.clean,
        *,
        normalize_fullwidth_colon: bool = True,
    ) -> None:
        self.cleaner = cleaner
        self.normalize_fullwidth_colon = normalize_fullwidth_colon

    def normalize(self, value: Any) -> str:
        text = self.cleaner(value)
        if self.normalize_fullwidth_colon:
            text = text.replace("：", ":")
        if not text:
            return ""
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?", text)
        if not match:
            return text
        hour = int(match.group(1))
        minute = int(match.group(2))
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return f"{hour:02d}:{minute:02d}"
        return text


class StrictEndOfDayTimeStrategy(TimeNormalizationStrategy):
    """Strict HH:MM normalization with 24:00 preserved in notes as 23:59."""

    def __init__(self, cleaner: Callable[[Any], str] = RowSupport.clean) -> None:
        self.cleaner = cleaner

    def normalize(
        self,
        value: Any,
        row: dict[str, str],
        source_name: str,
    ) -> str:
        text = self.cleaner(value)
        if not text:
            return ""
        if text == "24:00":
            RowSupport.append_note(row, f"原データ{source_name}=24:00")
            return "23:59"
        match = re.fullmatch(r"(\d{1,2}):(\d{2})", text)
        if not match:
            raise RuntimeError(f"unexpected {source_name}={text!r}")
        hour, minute = map(int, match.groups())
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            raise RuntimeError(f"invalid {source_name}={text!r}")
        return f"{hour:02d}:{minute:02d}"


class StrictHmTimeStrategy(TimeNormalizationStrategy):
    """Normalize strict HH:MM values and reject malformed input."""

    def __init__(self, cleaner: Callable[[Any], str] = RowSupport.clean) -> None:
        self.cleaner = cleaner

    def normalize(self, value: Any) -> str:
        text = self.cleaner(value).replace("：", ":")
        if not text:
            return ""
        match = re.fullmatch(r"(\d{1,2}):(\d{2})", text)
        if not match:
            raise RuntimeError(f"unexpected time value: {text!r}")
        hour, minute = map(int, match.groups())
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            raise RuntimeError(f"invalid time value: {text!r}")
        return f"{hour:02d}:{minute:02d}"


class LenientTwentyFourTimeStrategy(TimeNormalizationStrategy):
    """Normalize HH:MM through 24:00 and preserve unsupported text unchanged."""

    def __init__(self, cleaner: Callable[[Any], str] = RowSupport.clean) -> None:
        self.cleaner = cleaner

    def normalize(self, value: Any) -> str:
        text = self.cleaner(value)
        if not text:
            return ""
        parts = text.split(":")
        if len(parts) != 2:
            return text
        try:
            hour, minute = map(int, parts)
        except ValueError:
            return text
        if 0 <= hour <= 24 and 0 <= minute <= 59:
            if hour == 24 and minute != 0:
                return text
            return f"{hour:02d}:{minute:02d}"
        return text


class PreserveEndOfDayTimeStrategy(TimeNormalizationStrategy):
    """Normalize valid times and map 24:00 to 23:59 while preserving its source text."""

    def normalize(self, value: Any) -> tuple[str, str | None]:
        text = RowSupport.clean(value).replace("：", ":")
        if not text:
            return "", None
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
        if not match:
            return "", text
        hour = int(match.group(1))
        minute = int(match.group(2))
        second = match.group(3)
        if hour == 24 and minute == 0 and second in (None, "00"):
            return "23:59", text
        if 0 <= hour <= 23 and 0 <= minute <= 59 and second in (None, "00"):
            return f"{hour:02d}:{minute:02d}", None
        return "", text


class StableIdStrategy(ABC):
    @abstractmethod
    def generate(self, code: str, *parts: Any) -> str:
        raise NotImplementedError


class Sha256StableIdStrategy(StableIdStrategy):
    def __init__(
        self,
        digest_length: int = 16,
        template: str = "{code}-SRC-{digest}",
        include_code_in_material: bool = True,
    ) -> None:
        self.digest_length = digest_length
        self.template = template
        self.include_code_in_material = include_code_in_material

    def generate(self, code: str, *parts: Any) -> str:
        cleaned_parts = [RowSupport.clean(part) for part in parts]
        material_parts = [code, *cleaned_parts] if self.include_code_in_material else cleaned_parts
        material = "|".join(material_parts)
        digest = hashlib.sha256(material.encode("utf-8")).hexdigest()
        short_digest = digest[:self.digest_length].upper()
        return self.template.format(code=code, digest=short_digest)


class SourceReader(ABC):
    @abstractmethod
    def read(self, path: Path) -> SourceTable:
        raise NotImplementedError


class CsvDictSourceReader(SourceReader):
    def read(self, path: Path) -> SourceTable:
        encoding = detect_text_encoding(path)
        with path.open("r", encoding=encoding, newline="") as f:
            reader = csv.DictReader(f)
            header = list(reader.fieldnames or [])
            rows: list[dict[str, str]] = []
            for line_no, source_row in enumerate(reader, 2):
                if None in source_row:
                    raise RuntimeError(
                        f"{path.name}:{line_no}: overflow columns={source_row[None]!r}"
                    )
                row = {key: RowSupport.clean(value) for key, value in source_row.items()}
                if RowSupport.has_meaningful_value(row.values()):
                    rows.append(row)
        return SourceTable(header=header, rows=rows, source_format=encoding)


class XlsxSourceReader(SourceReader):
    def __init__(self, sheet_name: str) -> None:
        self.sheet_name = sheet_name

    @staticmethod
    def _trim_trailing_blank_headers(header: list[str]) -> list[str]:
        end = len(header)
        while end > 0 and not header[end - 1]:
            end -= 1
        return header[:end]

    def read(self, path: Path) -> SourceTable:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if self.sheet_name not in workbook.sheetnames:
                raise RuntimeError(
                    f"{path.name}: sheet {self.sheet_name!r} not found; "
                    f"sheets={workbook.sheetnames!r}"
                )
            worksheet = workbook[self.sheet_name]
            nonempty = [
                row
                for row in worksheet.iter_rows(values_only=True)
                if RowSupport.has_meaningful_value(row)
            ]
            if not nonempty:
                raise RuntimeError(f"{path.name}: empty sheet: {self.sheet_name}")

            raw_header = [RowSupport.clean(value) for value in nonempty[0]]
            header = self._trim_trailing_blank_headers(raw_header)
            rows: list[dict[str, str]] = []
            for values in nonempty[1:]:
                row = {
                    header[i]: RowSupport.clean(values[i] if i < len(values) else None)
                    for i in range(len(header))
                }
                if RowSupport.has_meaningful_value(row.values()):
                    rows.append(row)
            return SourceTable(header=header, rows=rows, source_format="xlsx")
        finally:
            workbook.close()


class StandardCsvWriter:
    def write(
        self,
        path: Path,
        schema: Sequence[str],
        rows: Sequence[dict[str, str]],
    ) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".csv.tmp")
        with tmp.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(schema), lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
        tmp.replace(path)
