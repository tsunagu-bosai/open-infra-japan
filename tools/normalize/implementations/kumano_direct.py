#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from tools.normalize.core.municipality import six_digit_municipality_code
from tools.normalize.core.normalization import (
    RowSupport,
    Sha256StableIdStrategy,
    StandardCsvWriter,
)
from tools.normalize.core.schema import read_schema_fields


ROOT = Path.cwd()

RAW = (
    ROOT
    / "data/raw/34-広島県/34307-熊野町/public-toilet/3a08.xlsx"
)

DEST = (
    ROOT
    / "data/normalized/34-広島県/34307-熊野町/"
    / "public-toilet/public-toilet.csv"
)

SCHEMA_PATH = ROOT / "schema/standard/public-toilet/schema.csv"

SHEET_NAME = "08.公衆トイレ一覧"
EXPECTED_ROWS = 5

EXPECTED_HEADER = (
    "市区町村コード",
    "NO",
    "市区町村名",
    "名称　必須",
    "名称_カナ　必須",
    "名称_英語",
    "住所　必須",
    "方書",
    "設置位置　必須",
    "緯度",
    "経度",
    "男性トイレ総数",
    "男性トイレ数（小便器）",
    "男性トイレ数（和式）",
    "男性トイレ数（洋式）",
    "女性トイレ総数",
    "女性トイレ数（和式）",
    "女性トイレ数（洋式）",
    "男女共用トイレ総数",
    "男女共用トイレ数（和式）",
    "男女共用トイレ数（洋式）",
    "多機能トイレ数　必須",
    "車椅子使用者用トイレ有無　必須",
    "乳幼児用設備設置トイレ有無　必須",
    "オストメイト設置トイレ有無　必須",
    "利用開始時間",
    "利用終了時間",
    "利用可能時間特記事項",
    "画像",
    "画像_ライセンス",
    "備考",
)

ID_STRATEGY = Sha256StableIdStrategy()
WRITER = StandardCsvWriter()


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_source() -> list[dict[str, str]]:
    if not RAW.exists():
        raise RuntimeError(f"missing source: {RAW}")

    wb = load_workbook(RAW, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise RuntimeError(
                f"{RAW.name}: missing sheet {SHEET_NAME!r}"
            )

        ws = wb[SHEET_NAME]

        header = tuple(
            text(v)
            for v in next(
                ws.iter_rows(
                    min_row=3,
                    max_row=3,
                    values_only=True,
                )
            )
        )

        if header != EXPECTED_HEADER:
            raise RuntimeError(
                f"{RAW.name}: unexpected header "
                f"columns={len(header)}"
            )

        rows = []

        for excel_row, values in enumerate(
            ws.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            row = {
                EXPECTED_HEADER[i]: text(value)
                for i, value in enumerate(values)
            }

            if not any(row.values()):
                continue

            # 4行目は配布テンプレートの記入例。
            if excel_row == 4:
                if (
                    row["名称　必須"] != "○○駅"
                    or row["市区町村コード"] != "034307"
                ):
                    raise RuntimeError(
                        f"{RAW.name}: expected sample row changed"
                    )
                continue

            rows.append(row)

    finally:
        wb.close()

    if len(rows) != EXPECTED_ROWS:
        raise RuntimeError(
            f"{RAW.name}: rows={len(rows)} "
            f"expected={EXPECTED_ROWS}"
        )

    return rows


def main() -> None:
    schema = read_schema_fields(SCHEMA_PATH)

    if len(schema) != 39:
        raise RuntimeError(
            f"schema columns={len(schema)} expected=39"
        )

    municipality_code = "34307"
    standard_code = six_digit_municipality_code(
        municipality_code
    )

    result: list[dict[str, str]] = []

    for src in read_source():
        if src["市区町村コード"] != "34307":
            raise RuntimeError(
                "unexpected municipality code: "
                f"{src['市区町村コード']!r}"
            )

        source_no = src["NO"]
        name = src["名称　必須"]
        address = src["住所　必須"]

        if not source_no or not name or not address:
            raise RuntimeError(
                "required field missing: "
                f"NO={source_no!r} "
                f"name={name!r} "
                f"address={address!r}"
            )

        row = {field: "" for field in schema}

        row["全国地方公共団体コード"] = standard_code
        row["所在地_全国地方公共団体コード"] = standard_code

        row["地方公共団体名"] = "熊野町"

        row["名称"] = name
        row["名称_カナ"] = src["名称_カナ　必須"]
        row["名称_英語"] = src["名称_英語"]

        row["所在地_連結表記"] = address
        row["所在地_都道府県"] = "広島県"
        row["所在地_市区町村"] = "安芸郡熊野町"
        row["建物名等(方書)"] = src["方書"]

        row["設置位置"] = src["設置位置　必須"]

        row["緯度"] = src["緯度"]
        row["経度"] = src["経度"]

        for field in (
            "男性トイレ総数",
            "男性トイレ数（小便器）",
            "男性トイレ数（和式）",
            "男性トイレ数（洋式）",
            "女性トイレ総数",
            "女性トイレ数（和式）",
            "女性トイレ数（洋式）",
            "男女共用トイレ総数",
            "男女共用トイレ数（和式）",
            "男女共用トイレ数（洋式）",
        ):
            row[field] = src[field]

        row["車椅子使用者用トイレ有無"] = (
            src["車椅子使用者用トイレ有無　必須"]
        )
        row["乳幼児用設備設置トイレ有無"] = (
            src["乳幼児用設備設置トイレ有無　必須"]
        )
        row["オストメイト設置トイレ有無"] = (
            src["オストメイト設置トイレ有無　必須"]
        )

        row["利用開始時間"] = src["利用開始時間"]
        row["利用終了時間"] = src["利用終了時間"]
        row["利用可能時間特記事項"] = (
            src["利用可能時間特記事項"]
        )

        row["画像"] = src["画像"]
        row["画像_ライセンス"] = src["画像_ライセンス"]

        row["ID"] = ID_STRATEGY.generate(
            municipality_code,
            "kumano-direct",
            source_no,
            name,
            address,
        )

        RowSupport.append_note(
            row,
            f"原データNO={source_no}",
        )

        multi = src["多機能トイレ数　必須"]
        if multi:
            RowSupport.append_note(
                row,
                f"原データ多機能トイレ数={multi}",
            )

        if src["備考"]:
            RowSupport.append_note(
                row,
                f"原データ備考={src['備考']}",
            )

        RowSupport.append_note(
            row,
            "出典=熊野町役場から直接提供（2026-10-01受領）",
        )

        result.append(row)

    ids = [row["ID"] for row in result]
    if len(ids) != len(set(ids)):
        raise RuntimeError("duplicate normalized IDs")

    if DEST.exists():
        raise RuntimeError(
            f"normalized already exists: {DEST}"
        )

    WRITER.write(DEST, schema, result)


if __name__ == "__main__":
    main()
