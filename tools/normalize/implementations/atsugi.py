#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from tools.normalize.core.normalization import (
    RowSupport,
    Sha256StableIdStrategy,
    StandardCsvWriter,
)
from tools.normalize.core.schema import read_schema_fields

ROOT = Path.cwd()
SCHEMA_PATH = ROOT / "schema/standard/public-toilet/schema.csv"

CODE = "14212"
CODE6 = "142123"
NAME = "厚木市"
PREF = "神奈川県"
PREF_DIR = "14-神奈川県"
MUN_DIR = "14212-厚木市"
SOURCE_FILE = "14212_public-toilet.xlsx"

SECTIONS = {
    "固定式トイレ一覧": "固定式",
    "移動式トイレ一覧": "移動式",
}


_ID_STRATEGY = Sha256StableIdStrategy(template="{code}-{digest}")
_WRITER = StandardCsvWriter()

def read_source(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if wb.sheetnames != ["Sheet1"]:
        raise RuntimeError(f"{path}: sheets={wb.sheetnames!r} expected=['Sheet1']")

    ws = wb["Sheet1"]
    current_section = ""
    rows: list[dict[str, str]] = []

    for row_no, values in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [RowSupport.clean(v) for v in values[:4]]
        if not any(cells):
            continue

        first = cells[0]

        if first in SECTIONS:
            current_section = SECTIONS[first]
            continue

        if first == "No.":
            if cells != ["No.", "設置場所", "型式", "設置数"]:
                raise RuntimeError(f"{path}:{row_no}: unexpected header={cells!r}")
            continue

        if first.replace("　", "").replace(" ", "") == "合計":
            continue

        if not current_section:
            raise RuntimeError(
                f"{path}:{row_no}: data row before section: {cells!r}"
            )

        no, place, kind, count = cells
        if not no or not place:
            raise RuntimeError(f"{path}:{row_no}: invalid data row={cells!r}")

        rows.append(
            {
                "section": current_section,
                "no": no,
                "place": place,
                "kind": kind,
                "count": count,
            }
        )

    if not rows:
        raise RuntimeError(f"{path}: no data rows")

    return rows


def prepare() -> tuple[Path, list[str], list[dict[str, str]]]:
    schema = read_schema_fields(SCHEMA_PATH)
    if len(schema) != 39:
        raise RuntimeError(f"{SCHEMA_PATH}: schema columns={len(schema)} expected=39")

    source_path = (
        ROOT / "data/raw" / PREF_DIR / MUN_DIR / "public-toilet" / SOURCE_FILE
    )
    dest = (
        ROOT / "data/normalized" / PREF_DIR / MUN_DIR
        / "public-toilet" / "public-toilet.csv"
    )

    source_rows = read_source(source_path)
    rows: list[dict[str, str]] = []
    seen_ids: set[str] = set()

    for source in source_rows:
        row = {field: "" for field in schema}

        row["全国地方公共団体コード"] = CODE6
        row["所在地_全国地方公共団体コード"] = CODE6
        row["地方公共団体名"] = NAME

        # 元データに住所・座標はなく、場所名のみ。
        # 位置を推測補完せず、場所名を名称・設置位置として保持する。
        row["名称"] = source["place"]
        row["所在地_都道府県"] = PREF
        row["所在地_市区町村"] = NAME
        row["設置位置"] = source["place"]

        row["ID"] = _ID_STRATEGY.generate(
            CODE,
            source["section"],
            source["no"],
            source["place"],
        )
        if row["ID"] in seen_ids:
            raise RuntimeError(f"duplicate generated ID={row['ID']}")
        seen_ids.add(row["ID"])

        # 型式・設置数は便器内訳等を含む独自表現のため、
        # 標準数値列へ推測変換せず原文を保存する。
        RowSupport.append_labeled_note(row, "原データ区分", source["section"])
        RowSupport.append_labeled_note(row, "原データNo.", source["no"])
        RowSupport.append_labeled_note(row, "原データ型式", source["kind"])
        RowSupport.append_labeled_note(row, "原データ設置数", source["count"])

        rows.append(row)

    return dest, schema, rows



def main() -> None:
    dest, schema, rows = prepare()
    print(f"preflight OK: {CODE} {NAME} rows={len(rows)}")
    _WRITER.write(dest, schema, rows)
    print(f"OK {CODE} {NAME}: rows={len(rows)} source={SOURCE_FILE}")


if __name__ == "__main__":
    main()
