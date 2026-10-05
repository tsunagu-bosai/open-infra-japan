import csv
import re
from pathlib import Path
from openpyxl import load_workbook

from tools.normalize.core.normalization import RowSupport, StandardCsvWriter

ROOT = Path.cwd()
RAW = ROOT / "data/raw/34-広島県"
OUT = ROOT / "data/normalized/34-広島県"
SCHEMA = ROOT / "schema/standard/public-toilet/schema.csv"

with SCHEMA.open(encoding="utf-8-sig", newline="") as f:
    schema_rows = list(csv.DictReader(f))

FIELDS = [r["name"] for r in schema_rows]
PHYSICAL_TO_NAME = {
    r["physical_name"]: r["name"]
    for r in schema_rows
    if r.get("physical_name")
}

# 現行標準 -> physical name
CURRENT = {
    "全国地方公共団体コード": "localgov_code",
    "ID": "id",
    "地方公共団体名": "localgov_name",
    "名称": "name",
    "名称_カナ": "name_kana",
    "名称_英語": "name_english",
    "所在地_全国地方公共団体コード": "seat_localgov_code",
    "町字ID": "choaza_id",
    "所在地_連結表記": "seat_copulative_spell",
    "所在地_都道府県": "seat_prefecture",
    "所在地_市区町村": "seat_city",
    "所在地_町字": "seat_choaza",
    "所在地_番地以下": "seat_after_address",
    "建物名等(方書)": "buillding_name_etc",
    "設置位置": "set_position",
    "緯度": "latitude",
    "経度": "longitude",
    "高度の種別": "advanced_classification",
    "高度の値": "advanced_value",
    "男性トイレ総数": "male_wc_total",
    "男性トイレ数（小便器）": "male_wc_urinal",
    "男性トイレ数（和式）": "male_wc_jstyle",
    "男性トイレ数（洋式）": "male_wc_wstyle",
    "女性トイレ総数": "female_wc_total",
    "女性トイレ数（和式）": "female_wc_jstyle",
    "女性トイレ数（洋式）": "female_wc_wstyle",
    "男女共用トイレ総数": "unisex_wc_total",
    "男女共用トイレ数（和式）": "unisex_wc_jstyle",
    "男女共用トイレ数（洋式）": "unisex_wc_wstyle",
    "バリアフリートイレ数": "barrier_free_wc",
    "車椅子使用者用トイレ有無": "wheelchair_wc",
    "乳幼児用設備設置トイレ有無": "infant_wc",
    "オストメイト設置トイレ有無": "ostomate_wc",
    "利用開始時間": "available_start_time",
    "利用終了時間": "available_end_time",
    "利用可能時間特記事項": "available_time_note",
    "画像": "image",
    "画像_ライセンス": "image_licence",
    "備考": "note",
}

CURRENT = {
    source: PHYSICAL_TO_NAME[target]
    for source, target in CURRENT.items()
}

# 旧32列系。意味が明確に対応する項目だけ。
LEGACY = {
    "NO": "id",
    "No": "id",
    "名称": "name",
    "名称_カナ": "name_kana",
    "名称_英語": "name_english",
    "住所": "seat_copulative_spell",
    "方書": "buillding_name_etc",
    "設置位置": "set_position",
    "緯度": "latitude",
    "経度": "longitude",
    "男性トイレ総数": "male_wc_total",
    "男性トイレ数（小便器）": "male_wc_urinal",
    "男性トイレ数（和式）": "male_wc_jstyle",
    "男性トイレ数（洋式）": "male_wc_wstyle",
    "女性トイレ総数": "female_wc_total",
    "女性トイレ数（和式）": "female_wc_jstyle",
    "女性トイレ数（洋式）": "female_wc_wstyle",
    "男女共用トイレ総数": "unisex_wc_total",
    "男女共用トイレ数（和式）": "unisex_wc_jstyle",
    "男女共用トイレ数（洋式）": "unisex_wc_wstyle",
    "車椅子使用者用トイレ有無": "wheelchair_wc",
    "車いす使用者用トイレ有無": "wheelchair_wc",
    "乳幼児用設備設置トイレ有無": "infant_wc",
    "オストメイト設置トイレ有無": "ostomate_wc",
    "利用開始時間": "available_start_time",
    "利用終了時間": "available_end_time",
    "利用可能時間特記事項": "available_time_note",
    "画像": "image",
    "画像_ライセンス": "image_licence",
    "備考": "note",
    "備　　　考": "note",
}

LEGACY = {
    source: PHYSICAL_TO_NAME[target]
    for source, target in LEGACY.items()
}

def s(v):
    return "" if v is None else str(v).strip()

def clean_multiline_text(value):
    text = s(value)
    if not text:
        return ""
    return "\n".join(line.rstrip() for line in text.splitlines())


def one_flag(value):
    value = s(value)
    if value == "1":
        return "有"
    if value in {"-1", "0"}:
        return "無"
    return ""


def circle_flag(value):
    value = s(value)
    if value in {"○", "◯"}:
        return "有"
    if value == "×":
        return "無"
    return ""


def normalize_special_time(value):
    value = s(value).replace("：", ":")

    if not value:
        return ""

    # openpyxl の time -> "08:30:00" 等
    parts = value.split(":")
    if len(parts) == 3 and parts[2] == "00":
        value = ":".join(parts[:2])

    parts = value.split(":")
    if len(parts) == 2:
        try:
            hour = int(parts[0])
            minute = int(parts[1])
            if 0 <= hour <= 23 and 0 <= minute <= 59:
                return f"{hour:02d}:{minute:02d}"
        except ValueError:
            pass

    # XLSX XML由来などで時刻がExcel小数になる場合にも対応
    try:
        fraction = float(value)
    except ValueError:
        return value

    if 0 <= fraction < 1:
        total_minutes = round(fraction * 24 * 60)
        hour, minute = divmod(total_minutes, 60)
        return f"{hour:02d}:{minute:02d}"

    return value


def extract_kure_overview(overview):
    overview = s(overview)

    result = {
        "設置位置": "",
        "利用開始時間": "",
        "利用終了時間": "",
        "車椅子使用者用トイレ有無": "",
    }

    m = re.search(
        r"設置位置[：:]\s*(.*?)(?=\s+(?:利用可能日|利用可能時間|休館日)[：:]|$)",
        overview,
    )
    if m:
        result["設置位置"] = m.group(1).strip()

    m = re.search(
        r"利用可能時間[：:]\s*"
        r"(\d{1,2})[：:](\d{2})\s*[〜～~-]\s*"
        r"(\d{1,2})[：:](\d{2})",
        overview,
    )
    if m:
        result["利用開始時間"] = f"{int(m.group(1)):02d}:{m.group(2)}"
        result["利用終了時間"] = f"{int(m.group(3)):02d}:{m.group(4)}"

    m = re.search(r"車イス対応トイレ[：:]\s*([○◯×])", overview)
    if m:
        result["車椅子使用者用トイレ有無"] = circle_flag(m.group(1))

    return result

def find_header_index(rows):
    header_markers = {
        "名称",
        "施設名",
        "地点名",
    }

    for i, row in enumerate(rows):
        values = {s(x) for x in row}
        if values & header_markers:
            return i

    raise RuntimeError("header row not found")

def read_csv(path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    else:
        raise RuntimeError(f"encoding: {path}")

    rows = list(csv.reader(text.splitlines()))

    # タイトル行付きにも対応。
    hi = find_header_index(rows)
    header = [s(x) for x in rows[hi]]

    return [
        {header[i]: s(v) for i, v in enumerate(r)
         if i < len(header) and header[i]}
        for r in rows[hi + 1:]
        if any(s(v) for v in r)
    ]

def read_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hi = find_header_index(rows)
    header = [s(x) for x in rows[hi]]

    return [
        {header[i]: s(v) for i, v in enumerate(r)
         if i < len(header) and header[i]}
        for r in rows[hi + 1:]
        if any(s(v) for v in r)
    ]

TARGETS = {
    "34202-呉市",
    "34204-三原市",
    "34208-府中市",
    "34210-庄原市",
    "34212-東広島市",
    "34304-海田町",
}

def main() -> None:
    total = 0

    for src in sorted(RAW.glob("*/public-toilet/*")):
        if src.parent.parent.name not in TARGETS:
            continue
        if src.suffix.lower() not in (".csv", ".xlsx"):
            continue

        municipality = src.parents[1].name

        rows = read_csv(src) if src.suffix.lower() == ".csv" else read_xlsx(src)
        code = municipality.split("-", 1)[0]

        is_current = rows and "全国地方公共団体コード" in rows[0]
        mapping = CURRENT if is_current else LEGACY

        output = []

        for r in rows:
            # ◎/○/△ のメタデータ行。
            # 実施設名ではないことを明示的に判定。
            name = s(r.get("名称"))
            if name in {"◎", "○", "△"}:
                continue

            # IDだけ入った実質空レコードは施設データではないため除外。
            # 「名称が空」という理由だけでは除外しない。
            meaningful = [
                v for k, v in r.items()
                if k not in {
                    "ID", "NO", "No",
                    "全国地方公共団体コード",
                    "都道府県コード又は市区町村コード",
                    "市区町村コード",
                    "地方公共団体名",
                    "都道府県名",
                    "市区町村名",
                }
                and s(v)
            ]
            if not meaningful:
                continue

            out = {field: "" for field in FIELDS}

            if code == "34202":
                # 呉市バリアフリーマップのうちトイレ行だけ。
                if s(r.get("区分")) != "toilet":
                    continue

                out["全国地方公共団体コード"] = "342025"
                out["地方公共団体名"] = "広島県呉市"
                out["名称"] = s(r.get("施設名"))
                out["所在地_連結表記"] = s(r.get("住所"))
                out["所在地_都道府県"] = "広島県"
                out["所在地_市区町村"] = "呉市"
                out["緯度"] = s(r.get("緯度"))
                out["経度"] = s(r.get("経度"))

                overview = s(r.get("概要"))
                parsed = extract_kure_overview(overview)

                out["設置位置"] = parsed["設置位置"]
                out["利用開始時間"] = parsed["利用開始時間"]
                out["利用終了時間"] = parsed["利用終了時間"]

                out["車椅子使用者用トイレ有無"] = (
                    parsed["車椅子使用者用トイレ有無"]
                )
                out["乳幼児用設備設置トイレ有無"] = circle_flag(
                    r.get("ベビーベッドはある？")
                )
                out["オストメイト設置トイレ有無"] = circle_flag(
                    r.get("オストメイトはある？")
                )

                # 原データの詳細を失わない。
                out["備考"] = overview

            elif code == "34208":
                # びんご府中おもてなしトイレ
                out["全国地方公共団体コード"] = s(
                    r.get("都道府県コード又は市区町村コード")
                )
                out["ID"] = s(r.get("NO"))
                out["地方公共団体名"] = "広島県府中市"
                out["名称"] = s(r.get("名称"))
                out["所在地_連結表記"] = s(r.get("住所"))
                out["所在地_都道府県"] = "広島県"
                out["所在地_市区町村"] = "府中市"
                out["建物名等(方書)"] = s(r.get("方書"))
                out["緯度"] = s(r.get("位置情報_緯度"))
                out["経度"] = s(r.get("位置情報_経度"))

                # 意味が明確なものだけ標準列へ移す。
                out["男性トイレ数（小便器）"] = s(
                    r.get("トイレ_男_小")
                )
                out["女性トイレ総数"] = s(
                    r.get("トイレ_女")
                )

                # 空欄を「無」と推測しない。
                if s(r.get("トイレ_車椅子")):
                    out["車椅子使用者用トイレ有無"] = (
                        "有" if s(r["トイレ_車椅子"]) != "0" else "無"
                    )

                if s(r.get("トイレ_乳幼児用設備設置")):
                    out["乳幼児用設備設置トイレ有無"] = (
                        "有"
                        if s(r["トイレ_乳幼児用設備設置"]) != "0"
                        else "無"
                    )

                available = s(r.get("利用時間"))
                if available == "終日":
                    out["利用開始時間"] = "00:00"
                    out["利用終了時間"] = "23:59"
                    out["利用可能時間特記事項"] = (
                        "原データ利用時間=終日"
                    )
                elif "-" in available:
                    start, end = available.split("-", 1)
                    out["利用開始時間"] = normalize_special_time(start)
                    out["利用終了時間"] = normalize_special_time(end)

                # 標準列へ直接落とせない値は保持。
                for source_field in (
                    "トイレ_男_大",
                    "トイレ_男_子_小",
                    "トイレ_女_子",
                    "トイレ_多目的",
                ):
                    value = s(r.get(source_field))
                    if value:
                        RowSupport.append_note(
                            out,
                            f"{source_field}:{value}",
                            separator="; ",
                        )

                if s(r.get("説明")):
                    RowSupport.append_note(out, "説明:" + s(r["説明"]), separator="; ")

                if s(r.get("備考")):
                    RowSupport.append_note(out, s(r["備考"]), separator="; ")

            elif code == "34212":
                # 東広島市の施設データから、公衆トイレかつ市内分のみ。
                if s(r.get("公衆トイレ")) != "1":
                    continue

                address = s(r.get("住所"))
                if "東広島市" not in address:
                    continue

                out["全国地方公共団体コード"] = "342122"
                out["ID"] = s(r.get("ID"))
                out["地方公共団体名"] = "広島県東広島市"
                out["名称"] = s(r.get("地点名"))
                out["所在地_連結表記"] = address
                out["所在地_都道府県"] = "広島県"
                out["所在地_市区町村"] = "東広島市"

                out["車椅子使用者用トイレ有無"] = one_flag(
                    r.get("車イス対応")
                )
                out["乳幼児用設備設置トイレ有無"] = one_flag(
                    r.get("おむつ交換設備")
                )
                out["オストメイト設置トイレ有無"] = one_flag(
                    r.get("オストメイト対応")
                )

                out["利用開始時間"] = normalize_special_time(
                    r.get("営業時間1開始")
                )
                out["利用終了時間"] = normalize_special_time(
                    r.get("営業時間1終了")
                )

                if s(r.get("メモ")):
                    out["備考"] = clean_multiline_text(r["メモ"])

                # 標準列に直接対応しない有用情報を保存。
                for source_field in (
                    "洋式",
                    "和式",
                    "男女別",
                    "バリアフリートイレ",
                    "大型ベッド",
                ):
                    value = s(r.get(source_field))
                    if value:
                        RowSupport.append_note(
                            out,
                            f"{source_field}:{value}",
                            separator="; ",
                        )

            else:
                # 三原市・庄原市・海田町など既存形式。
                for source_name, target_name in mapping.items():
                    if target_name in out and source_name in r:
                        out[target_name] = s(r[source_name])

                # 旧データは原典に明示された値だけを標準列へ移す。
                if not is_current:
                    out["地方公共団体名"] = s(
                        r.get("市区町村名") or r.get("都道府県名")
                    )
                    out["所在地_都道府県"] = s(r.get("都道府県名"))
                    out["所在地_市区町村"] = s(r.get("市区町村名"))

                    legacy_code = s(
                        r.get("都道府県コード又は市区町村コード")
                        or r.get("市区町村コード")
                    )

                    if len(legacy_code) == 6 and legacy_code.isdigit():
                        out["全国地方公共団体コード"] = legacy_code

            # 多機能トイレ数/有無はバリアフリートイレ数へ推測変換しない。
            extras = []

            if s(r.get("多機能トイレ数")):
                extras.append("多機能トイレ数:" + s(r["多機能トイレ数"]))

            if s(r.get("多機能トイレ有無")):
                extras.append("多機能トイレ有無:" + s(r["多機能トイレ有無"]))

            if extras:
                existing = out["備考"]
                out["備考"] = (
                    existing + ("; " if existing else "") + "; ".join(extras)
                )

            # 有無項目は標準語彙「有」「無」のみ採用する。
            # 0/1/2/3等は意味を推測せず、原典値を備考へ保存する。
            for source_field, target_field in (
                ("車椅子使用者用トイレ有無", "車椅子使用者用トイレ有無"),
                ("車いす使用者用トイレ有無", "車椅子使用者用トイレ有無"),
                ("乳幼児用設備設置トイレ有無", "乳幼児用設備設置トイレ有無"),
                ("オストメイト設置トイレ有無", "オストメイト設置トイレ有無"),
            ):
                raw_value = s(r.get(source_field))
                if raw_value and raw_value not in {"有", "無"}:
                    if out.get(target_field) == raw_value:
                        out[target_field] = ""

                    note = out.get("備考", "")
                    preserved = f"{source_field}:{raw_value}"
                    if preserved not in note:
                        out["備考"] = (
                            note + ("; " if note else "") + preserved
                        )

            # 明白な緯度・経度逆転のみ補正する。
            # 日本の緯度・経度の値域から逆転が一意に判断できる場合に限定。
            try:
                lat = float(out.get("緯度", ""))
                lon = float(out.get("経度", ""))
                if 120 <= lat <= 150 and 20 <= lon <= 50:
                    out["緯度"], out["経度"] = out["経度"], out["緯度"]
            except (ValueError, TypeError):
                pass

            # 時刻表記を意味を変えない範囲で HH:MM に統一。
            for time_field in ("利用開始時間", "利用終了時間"):
                value = s(out.get(time_field))

                parts = value.split(":")

                # HH:MM:00 -> HH:MM
                if len(parts) == 3 and parts[2] == "00":
                    parts = parts[:2]

                if len(parts) == 2:
                    try:
                        hour = int(parts[0])
                        minute = int(parts[1])
                        if 0 <= hour <= 24 and 0 <= minute <= 59:
                            out[time_field] = f"{hour:02d}:{minute:02d}"
                    except ValueError:
                        pass

            output.append(out)

        dest = OUT / municipality / "public-toilet/public-toilet.csv"
        StandardCsvWriter().write(dest, FIELDS, output)

        print(
            municipality,
            "source=" + str(len(rows)),
            "normalized=" + str(len(output)),
            "type=" + ("current39" if is_current else "legacy")
        )

        total += len(output)

    print("total normalized:", total)


if __name__ == "__main__":
    main()
