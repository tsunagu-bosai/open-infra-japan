#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from tools import normalize_public_toilet as common
from tools.normalize.core.normalization import (
    AllowedValueStrategy,
    RowSupport,
    StandardCsvWriter,
    StrictHmsTimeStrategy,
)

ROOT = Path(__file__).resolve().parents[3]
SCHEMA = ROOT / "schema/standard/public-toilet/schema.csv"
RAW_ROOT = ROOT / "data/raw/19-山梨県"
OUT_ROOT = ROOT / "data/normalized/19-山梨県"

TARGETS = {
    "19211": "笛吹市",
    "19365": "身延町",
}

FUEFUKI_SOURCE = RAW_ROOT / "19211-笛吹市/public-toilet/koukyoushisetsutoire.xlsx"
FUEFUKI_OUTPUT = OUT_ROOT / "19211-笛吹市/public-toilet/public-toilet.csv"

MINOBU_SOURCE = RAW_ROOT / "19365-身延町/public-toilet/2851.xlsx"
MINOBU_OUTPUT = OUT_ROOT / "19365-身延町/public-toilet/public-toilet.csv"

FUEFUKI_HEADER = [
    "都道府県コード\n又は市区町村コード", "NO", "都道府県名", "市区町村名",
    "名称", "名称_カナ", "名称_英語", "住所", "設置位置", "緯度", "経度",
    "男性トイレ総数", "男性トイレ数（小便器）", "男性トイレ数（和式）",
    "男性トイレ数（洋式）", "女性トイレ総数", "女性トイレ数（和式）",
    "女性トイレ数（洋式）", "男女共用トイレ総数",
    "男女共用トイレ数（和式）", "男女共用トイレ数（洋式）",
    "多機能トイレ数", "車椅子使用者用トイレ有無",
    "乳幼児用設備設置トイレ有無", "オストメイト設置トイレ有無",
    "利用開始時間", "利用終了時間", "利用可能時間特記事項",
    "画像", "画像_ライセンス", "備考",
]

MINOBU_HEADER = [
    "", "区分", "施設名", "住所", "方書", "電話", "ＦＡＸ",
    "担当課", "備考", "HP",
]



def read_schema() -> list[str]:
    with SCHEMA.open("r", encoding="utf-8-sig", newline="") as f:
        return [row["name"] for row in csv.DictReader(f)]



TIME_NORMALIZER = StrictHmsTimeStrategy()
BOOLEAN_NORMALIZER = AllowedValueStrategy(("", "有", "無"), cleaner=RowSupport.clean)


def validate_count(value: str, field: str) -> str:
    value = RowSupport.clean(value)
    if value and not value.isdigit():
        raise RuntimeError(f"{field}: non-numeric value {value!r}")
    return value



_DMS_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)?)°\s*(\d+(?:\.\d+)?)′\s*(\d+(?:\.\d+)?)″\s*$"
)


def dms_to_decimal(value: str, field: str) -> str:
    value = RowSupport.clean(value)
    if not value:
        return ""
    m = _DMS_RE.fullmatch(value)
    if not m:
        raise RuntimeError(f"{field}: unexpected DMS value {value!r}")
    deg, minute, second = map(float, m.groups())
    if not (0 <= minute < 60 and 0 <= second < 60):
        raise RuntimeError(f"{field}: invalid DMS value {value!r}")
    decimal = deg + minute / 60.0 + second / 3600.0
    return f"{decimal:.6f}"



def normalize_fuefuki(schema: list[str]):
    if FUEFUKI_OUTPUT.exists():
        raise RuntimeError(f"19211: normalized already exists: {FUEFUKI_OUTPUT}")

    wb = load_workbook(FUEFUKI_SOURCE, read_only=True, data_only=True)
    ws = wb["公衆トイレ一覧_フォーマット"]
    physical = [[RowSupport.clean(v) for v in row] for row in ws.iter_rows(values_only=True)]

    if not physical or physical[0] != FUEFUKI_HEADER:
        raise RuntimeError("19211: unexpected source header")

    source_rows = [row for row in physical[1:] if any(RowSupport.clean(v) for v in row)]
    if not source_rows:
        raise RuntimeError("19211: source contains no meaningful data rows")

    src_dicts = [dict(zip(FUEFUKI_HEADER, row)) for row in source_rows]
    raw_ids = [RowSupport.clean(r["NO"]) for r in src_dicts]
    counts = Counter(raw_ids)
    if counts["0000000023"] != 2:
        raise RuntimeError(
            f"19211: duplicate NO distribution changed: {counts['0000000023']}"
        )

    # Stable duplicate ranking independent of physical row order.
    duplicate_rank: dict[int, int] = {}
    grouped: defaultdict[str, list[tuple[tuple[str, ...], int]]] = defaultdict(list)
    for idx, src in enumerate(src_dicts):
        no = RowSupport.clean(src["NO"])
        if counts[no] > 1:
            key = (
                RowSupport.clean(src["設置位置"]),
                RowSupport.clean(src["名称"]),
                RowSupport.clean(src["住所"]),
                RowSupport.clean(src["緯度"]),
                RowSupport.clean(src["経度"]),
                RowSupport.clean(src["オストメイト設置トイレ有無"]),
            )
            grouped[no].append((key, idx))
    for no, items in grouped.items():
        for rank, (_, idx) in enumerate(sorted(items), start=1):
            duplicate_rank[idx] = rank

    count_fields = [
        "男性トイレ総数", "男性トイレ数（小便器）", "男性トイレ数（和式）",
        "男性トイレ数（洋式）", "女性トイレ総数", "女性トイレ数（和式）",
        "女性トイレ数（洋式）", "男女共用トイレ総数",
        "男女共用トイレ数（和式）", "男女共用トイレ数（洋式）",
    ]

    normalized = []
    seen_ids: set[str] = set()
    dms_converted = 0
    duplicate_ids_resolved = 0

    for idx, src in enumerate(src_dicts):
        if RowSupport.clean(src["都道府県コード\n又は市区町村コード"]) != "192112":
            raise RuntimeError("19211: unexpected municipality code")

        raw_id = RowSupport.clean(src["NO"])
        if not raw_id:
            raise RuntimeError("19211: blank NO")

        rank = duplicate_rank.get(idx)
        ident = raw_id if rank in (None, 1) else f"{raw_id}-{rank}"
        if ident in seen_ids:
            raise RuntimeError(f"19211: duplicate normalized ID: {ident}")
        seen_ids.add(ident)

        row = {field: "" for field in schema}
        row["全国地方公共団体コード"] = "192112"
        row["ID"] = ident
        row["地方公共団体名"] = "山梨県笛吹市"
        row["名称"] = RowSupport.clean(src["名称"])
        row["名称_カナ"] = RowSupport.clean(src["名称_カナ"])
        row["名称_英語"] = RowSupport.clean(src["名称_英語"])
        row["所在地_全国地方公共団体コード"] = "192112"
        row["所在地_連結表記"] = RowSupport.clean(src["住所"])
        row["所在地_都道府県"] = "山梨県"
        row["所在地_市区町村"] = "笛吹市"
        row["設置位置"] = RowSupport.clean(src["設置位置"])

        raw_lat = RowSupport.clean(src["緯度"])
        raw_lon = RowSupport.clean(src["経度"])
        if bool(raw_lat) != bool(raw_lon):
            raise RuntimeError(
                f"19211: latitude/longitude presence mismatch: {raw_lat!r}, {raw_lon!r}"
            )
        if raw_lat:
            row["緯度"] = dms_to_decimal(raw_lat, "緯度")
            row["経度"] = dms_to_decimal(raw_lon, "経度")
            RowSupport.append_note(row, f"原データ緯度={raw_lat}")
            RowSupport.append_note(row, f"原データ経度={raw_lon}")
            dms_converted += 1

        for field in count_fields:
            row[field] = validate_count(src[field], field)

        for field in (
            "車椅子使用者用トイレ有無",
            "乳幼児用設備設置トイレ有無",
            "オストメイト設置トイレ有無",
        ):
            row[field] = BOOLEAN_NORMALIZER(src[field], field)

        row["利用開始時間"] = TIME_NORMALIZER(src["利用開始時間"])
        row["利用終了時間"] = TIME_NORMALIZER(src["利用終了時間"])
        row["利用可能時間特記事項"] = RowSupport.clean(src["利用可能時間特記事項"])
        row["画像"] = RowSupport.clean(src["画像"])
        row["画像_ライセンス"] = RowSupport.clean(src["画像_ライセンス"])
        RowSupport.append_note(row, RowSupport.clean(src["備考"]))

        multi = RowSupport.clean(src["多機能トイレ数"])
        if multi:
            RowSupport.append_note(row, f"原データ多機能トイレ数={multi}")

        if counts[raw_id] > 1:
            RowSupport.append_note(row, f"原データNO重複={raw_id}")
            if rank and rank > 1:
                RowSupport.append_note(
                    row,
                    f"重複解消のためIDを{ident}に変更"
                )
                duplicate_ids_resolved += 1

        normalized.append(row)

    patches = common.load_patches("19-山梨県", "19211-笛吹市")
    applied, problems = common.apply_patches(normalized, patches, FUEFUKI_SOURCE)
    if problems:
        raise RuntimeError(
            "19211-笛吹市: patch適用失敗: " + " | ".join(map(str, problems))
        )

    StandardCsvWriter().write(FUEFUKI_OUTPUT, schema, normalized)
    return len(normalized), dms_converted, duplicate_ids_resolved, applied


def normalize_minobu(schema: list[str]):
    if MINOBU_OUTPUT.exists():
        raise RuntimeError(f"19365: normalized already exists: {MINOBU_OUTPUT}")

    wb = load_workbook(MINOBU_SOURCE, read_only=True, data_only=True)
    ws = wb["公衆トイレ"]
    physical = [[RowSupport.clean(v) for v in row] for row in ws.iter_rows(values_only=True)]

    if not physical or physical[0] != MINOBU_HEADER:
        raise RuntimeError(f"19365: unexpected source header: {physical[0]!r}")

    source_rows = [row for row in physical[1:] if any(RowSupport.clean(v) for v in row)]
    if not source_rows:
        raise RuntimeError("19365: source contains no meaningful data rows")

    normalized = []
    seen_ids: set[str] = set()

    for row_values in source_rows:
        source_no = RowSupport.clean(row_values[0])
        if not source_no.isdigit():
            raise RuntimeError(f"19365: invalid source ordinal: {source_no!r}")
        ident = f"19365-SRC-{int(source_no):04d}"
        if ident in seen_ids:
            raise RuntimeError(f"19365: duplicate ID: {ident}")
        seen_ids.add(ident)

        src = dict(zip(MINOBU_HEADER[1:], row_values[1:]))
        if RowSupport.clean(src["区分"]) != "公衆トイレ":
            raise RuntimeError(f"19365: unexpected 区分={src['区分']!r}")

        out = {field: "" for field in schema}
        out["全国地方公共団体コード"] = "193658"
        out["ID"] = ident
        out["地方公共団体名"] = "山梨県身延町"
        out["名称"] = RowSupport.clean(src["施設名"])
        out["所在地_全国地方公共団体コード"] = "193658"

        address = RowSupport.clean(src["住所"])
        out["所在地_連結表記"] = (
            address if address.startswith("山梨県") else f"山梨県{address}"
        )
        out["所在地_都道府県"] = "山梨県"
        out["所在地_市区町村"] = "身延町"
        out["建物名等(方書)"] = RowSupport.clean(src["方書"])
        out["備考"] = RowSupport.clean(src["備考"])

        RowSupport.append_note(out, f"原データ先頭無名列={source_no}")
        RowSupport.append_note(out, f"原データ区分={RowSupport.clean(src['区分'])}")
        if RowSupport.clean(src["担当課"]):
            RowSupport.append_note(out, f"担当課={RowSupport.clean(src['担当課'])}")
        if RowSupport.clean(src["電話"]):
            RowSupport.append_note(out, f"電話={RowSupport.clean(src['電話'])}")
        if RowSupport.clean(src["ＦＡＸ"]):
            RowSupport.append_note(out, f"FAX={RowSupport.clean(src['ＦＡＸ'])}")
        if RowSupport.clean(src["HP"]):
            RowSupport.append_note(out, f"HP={RowSupport.clean(src['HP'])}")

        normalized.append(out)

    patches = common.load_patches("19-山梨県", "19365-身延町")
    applied, problems = common.apply_patches(normalized, patches, MINOBU_SOURCE)
    if problems:
        raise RuntimeError(
            "19365-身延町: patch適用失敗: " + " | ".join(map(str, problems))
        )

    StandardCsvWriter().write(MINOBU_OUTPUT, schema, normalized)
    return len(normalized), applied


def main() -> None:
    schema = read_schema()
    prepared = []

    if "19211" in TARGETS:
        prepared.append(("19211", normalize_fuefuki(schema)))
    if "19365" in TARGETS:
        prepared.append(("19365", normalize_minobu(schema)))

    for code, result in prepared:
        if code == "19211":
            rows, dms, dup, patches = result
            print(
                f"OK 19211 笛吹市: rows={rows} dms_converted={dms} "
                f"duplicate_ids_resolved={dup} patches={patches}"
            )
        elif code == "19365":
            rows, patches = result
            print(
                f"OK 19365 身延町: rows={rows} source_ordinal_ids={rows} "
                f"patches={patches}"
            )


if __name__ == "__main__":
    main()
