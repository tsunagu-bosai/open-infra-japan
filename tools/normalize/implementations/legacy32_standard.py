#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

from tools.normalize.core.encoding import detect_text_encoding
from tools.normalize.core.municipality import six_digit_municipality_code
from tools.normalize.core.schema import read_schema_fields
from tools.normalize.core.normalization import (
    LenientHmsTimeStrategy,
    RowSupport,
    Sha256StableIdStrategy,
    StandardCsvWriter,
)

ROOT = Path.cwd()
RAW = ROOT / "data/raw"
OUT = ROOT / "data/normalized"
SCHEMA_PATH = ROOT / "schema/standard/public-toilet/schema.csv"

TARGETS = {
    "11230": ("新座市", "11-埼玉県", "11230_public-toilet.csv"),
    "01231": ("恵庭市", "01-北海道", "012319_public_toilet-8002.csv"),
    "01304": ("新篠津村", "01-北海道", "013048_public_toilet-717.csv"),
    "01425": ("上砂川町", "01-北海道", "01425_public-toilet.csv"),
    "01465": ("剣淵町", "01-北海道", "01465_public-toilet.xlsx"),
    "01604": ("新冠町", "01-北海道", "01604_public-toilet.xlsx"),
    "01647": ("足寄町", "01-北海道", "public_toilet-1962.csv"),
    "01691": ("別海町", "01-北海道", "016918_public_toilet-2330.csv"),
    "03201": ("盛岡市", "03-岩手県", "public-toilet.csv"),
    "03202": ("宮古市", "03-岩手県", "public-toilet.csv"),
    "03215": ("奥州市", "03-岩手県", "03215_public-toilet.csv"),
    "04212": ("登米市", "04-宮城県", "04212_public-toilet.csv"),
    "07207": ("須賀川市", "07-福島県", "07207_public-toilet.csv"),
    "08205": ("石岡市", "08-茨城県", "08205_public-toilet.xlsx"),
    "08235": ("つくばみらい市", "08-茨城県", "08235_public-toilet.xlsx"),
    "09214": ("さくら市", "09-栃木県", "61bb27f8ceabf.csv"),
    "12220": ("流山市", "12-千葉県", "public_toilet.csv"),
    "12329": ("栄町", "12-千葉県", "public_toilet.csv"),
    "13116": ("豊島区", "13-東京都", "13116_public-toilet.csv"),
    "13221": ("清瀬市", "13-東京都", "13221_public-toilet.csv"),
    "14207": ("茅ヶ崎市", "14-神奈川県", "public_toilet.csv"),
    "14362": ("大井町", "14-神奈川県", "public_toilet.csv"),
    "16204": ("魚津市", "16-富山県", "064927.csv"),
    "17202": ("七尾市", "17-石川県", "08_koshutoire_2.csv"),
    "17206": ("加賀市", "17-石川県", "172065_public_toilet_202206.csv"),
    "17207": ("羽咋市", "17-石川県", "toilet_data.csv"),
    "17407": ("中能登町", "17-石川県", "174076_public_toilet_utf-8.csv"),
    "20214": ("茅野市", "20-長野県", "20214_public-toilet.csv"),
    "20403": ("高森町", "20-長野県", "204030_8_public_toilet.csv"),
    "21203": ("高山市", "21-岐阜県", "21203_public-toilet.csv"),
    "23209": ("碧南市", "23-愛知県", "232092_public_toilet.csv"),
    "23214": ("蒲郡市", "23-愛知県", "23214_public-toilet.csv"),
    "23215": ("犬山市", "23-愛知県", "toilet20251219.csv"),
    "23228": ("岩倉市", "23-愛知県", "232289_public_toilet.csv"),
    "23231": ("田原市", "23-愛知県", "23231_public-toilet.csv"),
    "25202": ("彦根市", "25-滋賀県", "25202_public-toilet.csv"),
    "25212": ("高島市", "25-滋賀県", "25212_public-toilet.csv"),
    "26201": ("福知山市", "26-京都府", "26201_public-toilet.csv"),
    "27208": ("貝塚市", "27-大阪府", "27208_public-toilet.csv"),
    "27209": ("守口市", "27-大阪府", "27209_public-toilet.xlsx"),
    "27223": ("門真市", "27-大阪府", "27223_public-toilet.csv"),
    "27232": ("阪南市", "27-大阪府", "27232_public-toilet.csv"),
    "28204": ("西宮市", "28-兵庫県", "xxxxxx_public_toilet.csv"),
    "29203": ("大和郡山市", "29-奈良県", "292036_public_toilet.csv"),
    "29208": ("御所市", "29-奈良県", "292087_public_toilet.csv"),
    "29453": ("東吉野村", "29-奈良県", "294535_public_toilet.csv"),
    "32207": ("江津市", "32-島根県", "32207_public-toilet.csv"),
    "36202": ("鳴門市", "36-徳島県", "36202_public-toilet.csv"),
    "38215": ("東温市", "38-愛媛県", "1979.csv"),
    "36387": ("美波町", "36-徳島県", "36387_public-toilet.csv"),
    "38402": ("砥部町", "38-愛媛県", "38402_public-toilet.csv"),
    "40218": ("春日市", "40-福岡県", "40218_public-toilet.csv"),
    "42203": ("島原市", "42-長崎県", "42203_public-toilet.csv"),
    "42308": ("時津町", "42-長崎県", "42308_public-toilet.csv"),
    "43348": ("美里町", "43-熊本県", "43348_public-toilet.csv"),
    "44208": ("竹田市", "44-大分県", "44208_public-toilet.csv"),
    "46214": ("垂水市", "46-鹿児島県", "46214_public-toilet.csv"),
    "46225": ("姶良市", "46-鹿児島県", "46225_public-toilet.csv"),
    "46492": ("肝付町", "46-鹿児島県", "46492_public-toilet.csv"),
    "46502": ("南種子町", "46-鹿児島県", "46502_public-toilet.csv"),
    "46535": ("与論町", "46-鹿児島県", "465356_public_toilet.csv"),
    "46223": ("南九州市", "46-鹿児島県", "46223_public-toilet.xlsx"),
    "47201": ("那覇市", "47-沖縄県", "47201_public-toilet.csv"),
    "47302": ("大宜味村", "47-沖縄県", "47302_public-toilet.csv"),
    "47357": ("南大東村", "47-沖縄県", "47357_public-toilet.csv"),
}

EXPECTED_MISSING = {
    "全国地方公共団体コード",
    "ID",
    "地方公共団体名",
    "所在地_全国地方公共団体コード",
    "町字ID",
    "所在地_連結表記",
    "所在地_都道府県",
    "所在地_市区町村",
    "所在地_町字",
    "所在地_番地以下",
    "建物名等(方書)",
    "高度の種別",
    "高度の値",
    "バリアフリートイレ数",
}

LEGACY_REQUIRED = {
    "NO", "都道府県名", "市区町村名", "住所", "方書", "多機能トイレ数"
}

CODE_VARIANTS = (
    "都道府県コード又は市区町村コード",
    "市区町村コード",
)


PRESERVE_EXTRA_COLUMNS = {
    "03201": ("所管課",),
}



TIME_NORMALIZER = LenientHmsTimeStrategy()


# BEGIN TAKAMORI LEGACY ID MAP
# 20403 高森町は元データのNOが全件空欄。
# 既存公開済み正規化CSVのIDを維持するため、
# 名称・住所・緯度・経度の安定キーから従来IDを再現する。
TAKAMORI_LEGACY_IDS = {('親水公園', '長野県下伊那郡高森町下市田2987番地先', '35.546556', '137.887444'): 'GEN-204030-0aa343631cca',
 ('お祭り広場', '長野県下伊那郡高森町吉田475番地41', '35.551065', '137.886614'): 'GEN-204030-a9ab197373a3',
 ('中央公園', '長野県下伊那郡高森町吉田2151番地6', '35.550288', '137.8834'): 'GEN-204030-b23cec275048',
 ('大丸山公園', '長野県下伊那郡高森町下市田2407番地1', '35.553217', '137.871599'): 'GEN-204030-c6da200246c6',
 ('城山公園', '長野県下伊那郡高森町吉田700番地4', '35.558302', '137.873156'): 'GEN-204030-28e1b4657214',
 ('天白公園', '長野県下伊那郡高森町牛牧2785番地8', '35.561566', '137.854156'): 'GEN-204030-f03eed48d651',
 ('天白幼児公園', '長野県下伊那郡高森町牛牧2785番地3', '35.561136', '137.855867'): 'GEN-204030-83379ecd3652',
 ('自然公園', '長野県下伊那郡高森町山吹2369番地6', '35.583703', '137.863874'): 'GEN-204030-ca6c5e9c05fb',
 ('山吹ミニゴルフ場', '長野県下伊那郡高森町山吹7251番地1', '35.586158', '137.871236'): 'GEN-204030-6250e7fcbc95',
 ('やまぶき公園', '長野県下伊那郡高森町山吹3624番地', '35.581236', '137.888577'): 'GEN-204030-9c29a7d72239',
 ('丸山公園', '長野県下伊那郡高森町山吹431番地3', '35.569744', '137.882186'): 'GEN-204030-8e52dc1700c2',
 ('吉田東公園', '長野県下伊那郡高森町吉田2321番地13', '35.55825', '137.889375'): 'GEN-204030-659a3e6b5021',
 ('牛牧マレットゴルフ場', '長野県下伊那郡高森町牛牧1915番地', '35.560084', '137.846168'): 'GEN-204030-fb15092d107d',
 ('中央道バス停（西）', '長野県下伊那郡高森町牛牧2614番地1', '35.558258', '137.857987'): 'GEN-204030-e56d941ba2a4',
 ('中央道バス停（東）', '長野県下伊那郡高森町牛牧2607番地2', '35.558053', '137.858737'): 'GEN-204030-c3a9e93bd336',
 ('不動滝', '長野県下伊那郡高森町牛牧', '35.587934', '137.831773'): 'GEN-204030-a49f5ef1b2c5',
 ('山の寺キャンプ場（上）', '長野県下伊那郡高森町山吹2347番地', '35.592789', '137.857081'): 'GEN-204030-694d2cb0751a',
 ('山の寺キャンプ場（下）', '長野県下伊那郡高森町山吹', '35.591284', '137.857189'): 'GEN-204030-4554911f2861',
 ('シルバーグリーンランド', '長野県下伊那郡高森町牛牧407番地4', '35.544909', '137.857997'): 'GEN-204030-19c0a6f6944f',
 ('町民グラウンド', '長野県下伊那郡高森町下市田2352番地', '35.549656', '137.871203'): 'GEN-204030-e08f13012bbb',
 ('新田グラウンド', '長野県下伊那郡高森町山吹7703番地3', '35.590348', '137.88051'): 'GEN-204030-5f770081851e'}
# END TAKAMORI LEGACY ID MAP

STABLE_ID = Sha256StableIdStrategy()
WRITER = StandardCsvWriter()


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_text(value) -> str:
    return re.sub(r"[\r\n]+", "", _cell_text(value)).strip()


def _row_has_meaningful_value(values) -> bool:
    for value in values:
        text = _cell_text(value)
        if text and text.replace(",", "").strip():
            return True
    return False


def read_csv_physical(path: Path, enc: str):
    with path.open("r", encoding=enc, newline="") as f:
        reader = csv.reader(f)
        header_values = next(reader, None)
        if header_values is None:
            raise RuntimeError(f"empty CSV: {path}")
        header = [_header_text(value) for value in header_values]
        if len(set(header)) != len(header):
            raise RuntimeError(f"duplicate header after normalization: {path}")

        rows = []
        mismatches = []
        for line_no, values in enumerate(reader, 2):
            if not _row_has_meaningful_value(values):
                continue
            if len(values) != len(header):
                mismatches.append((line_no, len(values), len(header)))
                continue
            rows.append({
                header[i]: _cell_text(values[i])
                for i in range(len(header))
            })

    return header, rows, mismatches


def read_xlsx_physical(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        values_iter = worksheet.iter_rows(values_only=True)
        header_values = next(values_iter, None)
        if header_values is None:
            raise RuntimeError(f"empty XLSX: {path}")

        header = [_header_text(value) for value in header_values]
        if len(set(header)) != len(header):
            raise RuntimeError(f"duplicate header after normalization: {path}")
        rows = []
        mismatches = []

        for line_no, values in enumerate(values_iter, 2):
            values = list(values)
            if not _row_has_meaningful_value(values):
                continue
            if len(values) != len(header):
                mismatches.append((line_no, len(values), len(header)))
                continue
            rows.append({
                header[i]: _cell_text(values[i])
                for i in range(len(header))
            })

        return header, rows, mismatches
    finally:
        workbook.close()


def read_physical(path: Path):
    if path.suffix.lower() == ".xlsx":
        header, rows, mismatches = read_xlsx_physical(path)
        return header, rows, mismatches, "xlsx"

    enc = detect_text_encoding(path)
    header, rows, mismatches = read_csv_physical(path, enc)
    return header, rows, mismatches, enc


def prepare_one(code, name, pref, filename, schema):
    src = RAW / pref / f"{code}-{name}" / "public-toilet" / filename
    dest = OUT / pref / f"{code}-{name}" / "public-toilet/public-toilet.csv"

    if not src.exists():
        raise RuntimeError(f"{code}: raw missing: {src}")
    if dest.exists():
        raise RuntimeError(f"{code}: normalized already exists: {dest}")

    header, rows, mismatches, source_format = read_physical(src)

    if mismatches:
        raise RuntimeError(f"{code}: physical row mismatch: {mismatches[:5]!r}")
    if not rows:
        raise RuntimeError(f"{code}: no data rows after blank-row filtering")
    preserve_extras = PRESERVE_EXTRA_COLUMNS.get(code, ())
    expected_header_columns = 32 + len(preserve_extras)
    if len(header) != expected_header_columns:
        raise RuntimeError(
            f"{code}: header columns={len(header)} "
            f"expected={expected_header_columns}"
        )

    missing = {c for c in schema if c not in header}
    if missing != EXPECTED_MISSING:
        raise RuntimeError(
            f"{code}: unexpected missing columns: "
            f"actual={sorted(missing)!r}"
        )

    code_cols = [c for c in CODE_VARIANTS if c in header]
    if len(code_cols) != 1:
        raise RuntimeError(f"{code}: code column candidates={code_cols!r}")
    code_col = code_cols[0]

    extras = [c for c in header if c not in schema]
    expected_extras = (
        set(LEGACY_REQUIRED)
        | set(preserve_extras)
        | {code_col}
    )
    if set(extras) != expected_extras:
        raise RuntimeError(
            f"{code}: unexpected legacy extras: "
            f"actual={extras!r} expected={sorted(expected_extras)!r}"
        )

    expected_code6 = six_digit_municipality_code(code)

    out_rows = []
    altered_source_codes = 0
    preserved_multi = 0
    tome_longitude_corrections = 0
    hannan_time_corrections = 0
    kiyose_count_corrections = 0
    kiyose_coordinate_corrections = 0
    nanao_count_corrections = 0
    kaga_all_day_corrections = 0
    kaga_lights_out_corrections = 0
    sakae_all_day_corrections = 0
    chigasaki_all_day_corrections = 0
    yoron_coordinate_swaps = 0

    for line_no, srcrow in enumerate(rows, 2):
        row = {c: srcrow.get(c, "") for c in schema}

        raw_code = srcrow.get(code_col, "").replace(",", "").strip()
        if raw_code:
            normalized_raw_code = raw_code.zfill(6)
            same_municipality = (
                normalized_raw_code == expected_code6
                or (
                    len(normalized_raw_code) == 6
                    and normalized_raw_code[:5] == code
                )
                or raw_code == code
            )
            if not same_municipality:
                raise RuntimeError(
                    f"{code} line {line_no}: source municipality code "
                    f"{raw_code!r} does not match expected {expected_code6!r}"
                )
            if raw_code != expected_code6:
                RowSupport.append_note(row, f"原データ自治体コード={raw_code}")
                altered_source_codes += 1

        row["全国地方公共団体コード"] = expected_code6
        row["所在地_全国地方公共団体コード"] = expected_code6

        no = srcrow.get("NO", "")

        if not no and code in {"14207", "14362"}:
            no = STABLE_ID.generate(
                code,
                srcrow.get("名称", ""),
                srcrow.get("住所", ""),
                srcrow.get("設置位置", ""),
                srcrow.get("緯度", ""),
                srcrow.get("経度", ""),
            )

        if not no and code == "20403":
            takamori_key = (
                (srcrow.get("名称") or "").strip(),
                (srcrow.get("住所") or "").strip(),
                (srcrow.get("緯度") or "").strip(),
                (srcrow.get("経度") or "").strip(),
            )
            no = TAKAMORI_LEGACY_IDS.get(takamori_key, "")
            if not no:
                raise RuntimeError(
                    f"{code} line {line_no}: 高森町の既存IDを再現できません: "
                    f"{takamori_key!r}"
                )
        row["ID"] = no

        row["地方公共団体名"] = name
        row["所在地_都道府県"] = srcrow.get("都道府県名", "")
        row["所在地_市区町村"] = srcrow.get("市区町村名", "")
        row["所在地_連結表記"] = srcrow.get("住所", "")
        row["建物名等(方書)"] = srcrow.get("方書", "")

        multi = srcrow.get("多機能トイレ数", "")
        if multi:
            RowSupport.append_note(row, f"原データ多機能トイレ数={multi}")
            preserved_multi += 1

            if code == "23209":
                if not multi.isdigit():
                    raise RuntimeError(
                        f"碧南市 unexpected 多機能トイレ数={multi!r}"
                    )
                row["バリアフリートイレ数"] = multi

        for field in preserve_extras:
            value = srcrow.get(field, "")
            if value:
                RowSupport.append_note(row, f"原データ{field}={value}")

        if (
            code == "04212"
            and row.get("ID") == "0000000063"
            and row.get("名称") == "中田地蔵沼親水公園"
        ):
            if row.get("経度") != "141264748":
                raise RuntimeError(
                    f"登米市 unexpected longitude={row.get('経度')!r}"
                )
            RowSupport.append_note(row, "原データ経度=141264748")
            row["経度"] = "141.264748"
            tome_longitude_corrections += 1

        if code == "13221" and row.get("名称") == "市民活動センター":
            original_count = "2※内1個は故障中"
            if row.get("男性トイレ数（小便器）") != original_count:
                raise RuntimeError(
                    "清瀬市 unexpected urinal count="
                    f"{row.get('男性トイレ数（小便器）')!r}"
                )
            RowSupport.append_note(
                row,
                f"原データ男性トイレ数（小便器）={original_count}",
            )
            row["男性トイレ数（小便器）"] = "2"
            kiyose_count_corrections += 1

        if code == "13221" and row.get("名称") == "コミュニティプラザひまわり":
            original_lat = '''35°46'43.2"N'''
            original_lon = '''139°32'19.1"E'''

            if row.get("緯度") != original_lat:
                raise RuntimeError(
                    f"清瀬市 unexpected latitude={row.get('緯度')!r}"
                )
            if row.get("経度") != original_lon:
                raise RuntimeError(
                    f"清瀬市 unexpected longitude={row.get('経度')!r}"
                )

            RowSupport.append_note(row, f"原データ緯度={original_lat}")
            RowSupport.append_note(row, f"原データ経度={original_lon}")
            row["緯度"] = "35.778667"
            row["経度"] = "139.538639"
            kiyose_coordinate_corrections += 1

        if code == "46535":
            original_lat = (row.get("緯度") or "").strip()
            original_lon = (row.get("経度") or "").strip()

            try:
                lat_value = float(original_lat)
                lon_value = float(original_lon)
            except ValueError as exc:
                raise RuntimeError(
                    f"与論町 invalid coordinates: "
                    f"lat={original_lat!r} lon={original_lon!r}"
                ) from exc

            if not (120 <= lat_value <= 150 and 20 <= lon_value <= 50):
                raise RuntimeError(
                    f"与論町 unexpected coordinates: "
                    f"lat={original_lat!r} lon={original_lon!r}"
                )

            RowSupport.append_note(row, f"原データ緯度={original_lat}")
            RowSupport.append_note(row, f"原データ経度={original_lon}")

            row["緯度"] = original_lon
            row["経度"] = original_lat
            yoron_coordinate_swaps += 1

        if code == "27232":
            if row.get("利用開始時間") != "使用制限なし":
                raise RuntimeError(
                    "阪南市 unexpected 利用開始時間="
                    f"{row.get('利用開始時間')!r}"
                )
            note = "原データ利用開始時間=使用制限なし"
            current = (row.get("利用可能時間特記事項") or "").strip()
            if note not in current:
                row["利用可能時間特記事項"] = (
                    f"{current} / {note}" if current else note
                )
            row["利用開始時間"] = ""
            hannan_time_corrections += 1

        if code == "17202" and row.get("名称") == "長浦うるおい公園トイレ":
            original = "1(幼児用ﾄｲﾚ)"
            if row.get("女性トイレ数（和式）") != original:
                raise RuntimeError(
                    "七尾市 unexpected 女性トイレ数（和式）="
                    f"{row.get('女性トイレ数（和式）')!r}"
                )
            RowSupport.append_note(row, f"原データ女性トイレ数（和式）={original}")
            row["女性トイレ数（和式）"] = "1"
            nanao_count_corrections += 1

        if code == "17206":
            start = (row.get("利用開始時間") or "").strip()
            end = (row.get("利用終了時間") or "").strip()

            if start == "終日" and end == "終日":
                note = "原データ利用時間=終日"
                current = (row.get("利用可能時間特記事項") or "").strip()
                if note not in current:
                    row["利用可能時間特記事項"] = (
                        f"{current} / {note}" if current else note
                    )
                row["利用開始時間"] = "00:00"
                row["利用終了時間"] = "23:59"
                kaga_all_day_corrections += 1

            elif end == "22:00消灯":
                note = "原データ利用終了時間=22:00消灯"
                current = (row.get("利用可能時間特記事項") or "").strip()
                if note not in current:
                    row["利用可能時間特記事項"] = (
                        f"{current} / {note}" if current else note
                    )
                row["利用終了時間"] = "22:00"
                kaga_lights_out_corrections += 1


        if code == "12329":
            start = (row.get("利用開始時間") or "").strip()
            end = (row.get("利用終了時間") or "").strip()
            note = (row.get("利用可能時間特記事項") or "").strip()

            if start != "無" or end != "無":
                raise RuntimeError(
                    "栄町 unexpected 利用時間="
                    f"{start!r} - {end!r}"
                )

            if note not in {"24時間使用可能", "24時間利用可能"}:
                raise RuntimeError(
                    "栄町 unexpected 利用可能時間特記事項="
                    f"{note!r}"
                )

            RowSupport.append_note(row, "原データ利用開始時間=無")
            RowSupport.append_note(row, "原データ利用終了時間=無")
            row["利用開始時間"] = "00:00"
            row["利用終了時間"] = "23:59"
            sakae_all_day_corrections += 1

        if code == "14207":
            start = (row.get("利用開始時間") or "").strip()
            end = (row.get("利用終了時間") or "").strip()
            note = (row.get("利用可能時間特記事項") or "").strip()

            if start == "00:00" and end == "00:00":
                if note != "24時間利用可能":
                    raise RuntimeError(
                        "茅ヶ崎市 unexpected 利用可能時間特記事項="
                        f"{note!r}"
                    )

                RowSupport.append_note(row, "原データ利用終了時間=00:00")
                row["利用終了時間"] = "23:59"
                chigasaki_all_day_corrections += 1

        row["利用開始時間"] = TIME_NORMALIZER(row["利用開始時間"])
        row["利用終了時間"] = TIME_NORMALIZER(row["利用終了時間"])

        out_rows.append(row)

    if code == "46535" and yoron_coordinate_swaps != 17:
        raise RuntimeError(
            "与論町 coordinate swaps="
            f"{yoron_coordinate_swaps} expected=17"
        )

    if code == "04212" and tome_longitude_corrections != 1:
        raise RuntimeError(
            "登米市 longitude corrections="
            f"{tome_longitude_corrections} expected=1"
        )

    if code == "27232" and hannan_time_corrections != 5:
        raise RuntimeError(
            "阪南市 time corrections="
            f"{hannan_time_corrections} expected=5"
        )

    if code == "17202" and nanao_count_corrections != 1:
        raise RuntimeError(
            "七尾市 count corrections="
            f"{nanao_count_corrections} expected=1"
        )

    if code == "17206":
        if kaga_all_day_corrections != 4:
            raise RuntimeError(
                "加賀市 all-day corrections="
                f"{kaga_all_day_corrections} expected=4"
            )
        if kaga_lights_out_corrections != 13:
            raise RuntimeError(
                "加賀市 lights-out corrections="
                f"{kaga_lights_out_corrections} expected=13"
            )

    if code == "14207" and chigasaki_all_day_corrections != 3:
        raise RuntimeError(
            "茅ヶ崎市 all-day corrections="
            f"{chigasaki_all_day_corrections} expected=3"
        )

    if code == "12329" and sakae_all_day_corrections != 6:
        raise RuntimeError(
            "栄町 all-day corrections="
            f"{sakae_all_day_corrections} expected=6"
        )

    if code == "13221":
        if kiyose_count_corrections != 1:
            raise RuntimeError(
                "清瀬市 count corrections="
                f"{kiyose_count_corrections} expected=1"
            )
        if kiyose_coordinate_corrections != 1:
            raise RuntimeError(
                "清瀬市 coordinate corrections="
                f"{kiyose_coordinate_corrections} expected=1"
            )

    return {
        "code": code,
        "name": name,
        "src": src,
        "dest": dest,
        "source_format": source_format,
        "rows": out_rows,
        "altered_source_codes": altered_source_codes,
        "preserved_multi": preserved_multi,
    }


def main():
    schema = read_schema_fields(SCHEMA_PATH)
    if len(schema) != 39:
        raise RuntimeError(f"schema columns={len(schema)} expected=39")

    prepared = []
    for code, (name, pref, filename) in TARGETS.items():
        prepared.append(
            prepare_one(code, name, pref, filename, schema)
        )

    print(f"preflight OK: {len(prepared)} municipalities")

    for item in prepared:
        WRITER.write(item["dest"], schema, item["rows"])
        print(
            f"OK {item['code']} {item['name']}: "
            f"rows={len(item['rows'])} source={item['source_format']} "
            f"code_notes={item['altered_source_codes']} "
            f"multi_notes={item['preserved_multi']}"
        )


if __name__ == "__main__":
    main()
