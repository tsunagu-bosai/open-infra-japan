#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from tools import normalize_public_toilet as common
from tools.normalize.core.normalization import (
    AllowedValueStrategy,
    RowSupport,
    StandardCsvWriter,
    StrictHmsTimeStrategy,
)

ROOT = Path(__file__).resolve().parents[3]
SCHEMA = ROOT / "schema/standard/public-toilet/schema.csv"
RAW_ROOT = ROOT / "data/raw/19-山梨県"
OUT_ROOT = ROOT / "data/normalized/19-山梨県"

TARGETS = {"19425": "山中湖村"}

SOURCE = RAW_ROOT / "19425-山中湖村/public-toilet/r5_od_koushutoire.xlsx"
OUTPUT = OUT_ROOT / "19425-山中湖村/public-toilet/public-toilet.csv"

CODE6 = "194255"
SOURCE_CODE = "194425"

EXPECTED_HEADER = [
    "都道府県コード又は市区町村コード", "NO", "都道府県名", "市区町村名",
    "名称", "名称_カナ", "住所", "設置位置", "緯度", "経度",
    "男性トイレ総数", "男性トイレ数（小便器）", "男性トイレ数（和式）",
    "男性トイレ数（洋式）", "女性トイレ総数", "女性トイレ数（和式）",
    "女性トイレ数（洋式）", "男女共用トイレ総数",
    "男女共用トイレ数（和式）", "男女共用トイレ数（洋式）",
    "多機能トイレ数", "車椅子使用者用トイレ有無",
    "乳幼児用設備設置トイレ有無", "オストメイト設置トイレ有無",
    "利用開始時間", "利用終了時間", "利用可能時間特記事項",
]


def read_schema() -> list[str]:
    with SCHEMA.open("r", encoding="utf-8-sig", newline="") as f:
        return [row["name"] for row in csv.DictReader(f)]


TIME_NORMALIZER = StrictHmsTimeStrategy(cleaner=RowSupport.clean_cell)
BOOLEAN_NORMALIZER = AllowedValueStrategy(("", "有", "無"), cleaner=RowSupport.clean_cell)



def validate_count(value: str, field: str) -> str:
    value = RowSupport.clean_cell(value)
    if value and not value.isdigit():
        raise RuntimeError(f"{field}: non-numeric value {value!r}")
    return value


def main() -> None:
    if "19425" not in TARGETS:
        return
    if OUTPUT.exists():
        raise RuntimeError(f"19425: normalized already exists: {OUTPUT}")

    schema = read_schema()
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["公衆トイレ一覧"]
    physical = [[RowSupport.clean_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]

    if not physical or physical[0] != EXPECTED_HEADER:
        raise RuntimeError(f"unexpected header: {physical[0] if physical else None!r}")

    source_rows = [row for row in physical[1:] if any(RowSupport.clean_cell(v) for v in row)]
    if not source_rows:
        raise RuntimeError("no active source rows")

    normalized = []
    ids: set[str] = set()

    count_fields = [
        "男性トイレ総数", "男性トイレ数（小便器）", "男性トイレ数（和式）",
        "男性トイレ数（洋式）", "女性トイレ総数", "女性トイレ数（和式）",
        "女性トイレ数（洋式）", "男女共用トイレ総数",
        "男女共用トイレ数（和式）", "男女共用トイレ数（洋式）",
    ]

    for values in source_rows:
        src = dict(zip(EXPECTED_HEADER, values))
        if RowSupport.clean_cell(src["都道府県コード又は市区町村コード"]) != SOURCE_CODE:
            raise RuntimeError(
                "unexpected source municipality code: "
                f"{src['都道府県コード又は市区町村コード']!r}"
            )

        ident = RowSupport.clean_cell(src["NO"])
        if not ident or ident in ids:
            raise RuntimeError(f"blank/duplicate NO: {ident!r}")
        ids.add(ident)

        row = {field: "" for field in schema}
        row["全国地方公共団体コード"] = CODE6
        row["ID"] = ident
        row["地方公共団体名"] = "山梨県山中湖村"
        row["名称"] = RowSupport.clean_cell(src["名称"])
        row["名称_カナ"] = RowSupport.clean_cell(src["名称_カナ"])
        row["所在地_全国地方公共団体コード"] = CODE6

        address = RowSupport.clean_cell(src["住所"])
        row["所在地_連結表記"] = (
            address if address.startswith("山梨県") else f"山梨県{address}"
        )
        row["所在地_都道府県"] = "山梨県"
        row["所在地_市区町村"] = "山中湖村"
        row["設置位置"] = RowSupport.clean_cell(src["設置位置"])

        lat = RowSupport.clean_cell(src["緯度"])
        lon = RowSupport.clean_cell(src["経度"])
        try:
            lat_num, lon_num = float(lat), float(lon)
        except ValueError as exc:
            raise RuntimeError(f"invalid coordinate: {lat!r}, {lon!r}") from exc
        if not (-90 <= lat_num <= 90 and -180 <= lon_num <= 180):
            raise RuntimeError(f"out-of-range coordinate: {lat!r}, {lon!r}")
        row["緯度"], row["経度"] = lat, lon

        for field in count_fields:
            row[field] = validate_count(src[field], field)

        for field in (
            "車椅子使用者用トイレ有無",
            "乳幼児用設備設置トイレ有無",
            "オストメイト設置トイレ有無",
        ):
            row[field] = BOOLEAN_NORMALIZER(src[field], field)

        row["利用開始時間"] = TIME_NORMALIZER(src["利用開始時間"])
        row["利用終了時間"] = TIME_NORMALIZER(src["利用終了時間"])
        row["利用可能時間特記事項"] = RowSupport.clean_cell(src["利用可能時間特記事項"])

        multi = RowSupport.clean_cell(src["多機能トイレ数"])
        if multi:
            RowSupport.append_note(row, f"原データ多機能トイレ数={multi}")

        RowSupport.append_note(row, f"原データ自治体コード={SOURCE_CODE}")
        normalized.append(row)

    patches = common.load_patches("19-山梨県", "19425-山中湖村")
    applied, problems = common.apply_patches(normalized, patches, SOURCE)
    if problems:
        raise RuntimeError(
            "19425-山中湖村: patch適用失敗: " + " | ".join(map(str, problems))
        )

    StandardCsvWriter().write(OUTPUT, schema, normalized)

    print(
        f"OK 19425 山中湖村: rows={len(normalized)} code={CODE6} "
        f"source_code_corrected={len(normalized)} patches={applied}"
    )


if __name__ == "__main__":
    main()
