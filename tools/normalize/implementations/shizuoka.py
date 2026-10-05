import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
RAW = ROOT / "data/raw/22-静岡県"
OUT = ROOT / "data/normalized/22-静岡県"
SCHEMA = ROOT / "schema/standard/public-toilet/schema.csv"

with SCHEMA.open(encoding="utf-8-sig", newline="") as f:
    schema_rows = list(csv.DictReader(f))

FIELDS = [r["name"] for r in schema_rows]
PHYSICAL_TO_NAME = {
    r["physical_name"]: r["name"]
    for r in schema_rows
    if r.get("physical_name")
}

# 現行標準 -> physical name
CURRENT = {
    "全国地方公共団体コード": "localgov_code",
    "ID": "id",
    "地方公共団体名": "localgov_name",
    "名称": "name",
    "名称_カナ": "name_kana",
    "名称_英語": "name_english",
    "所在地_全国地方公共団体コード": "seat_localgov_code",
    "町字ID": "choaza_id",
    "所在地_連結表記": "seat_copulative_spell",
    "所在地_都道府県": "seat_prefecture",
    "所在地_市区町村": "seat_city",
    "所在地_町字": "seat_choaza",
    "所在地_番地以下": "seat_after_address",
    "建物名等(方書)": "buillding_name_etc",
    "設置位置": "set_position",
    "緯度": "latitude",
    "経度": "longitude",
    "高度の種別": "advanced_classification",
    "高度の値": "advanced_value",
    "男性トイレ総数": "male_wc_total",
    "男性トイレ数（小便器）": "male_wc_urinal",
    "男性トイレ数（和式）": "male_wc_jstyle",
    "男性トイレ数（洋式）": "male_wc_wstyle",
    "女性トイレ総数": "female_wc_total",
    "女性トイレ数（和式）": "female_wc_jstyle",
    "女性トイレ数（洋式）": "female_wc_wstyle",
    "男女共用トイレ総数": "unisex_wc_total",
    "男女共用トイレ数（和式）": "unisex_wc_jstyle",
    "男女共用トイレ数（洋式）": "unisex_wc_wstyle",
    "バリアフリートイレ数": "barrier_free_wc",
    "車椅子使用者用トイレ有無": "wheelchair_wc",
    "乳幼児用設備設置トイレ有無": "infant_wc",
    "オストメイト設置トイレ有無": "ostomate_wc",
    "利用開始時間": "available_start_time",
    "利用終了時間": "available_end_time",
    "利用可能時間特記事項": "available_time_note",
    "画像": "image",
    "画像_ライセンス": "image_licence",
    "備考": "note",
}

CURRENT = {
    source: PHYSICAL_TO_NAME[target]
    for source, target in CURRENT.items()
}

# 旧32列系。意味が明確に対応する項目だけ。
LEGACY = {
    "NO": "id",
    "No": "id",
    "名称": "name",
    "名称_カナ": "name_kana",
    "名称_英語": "name_english",
    "住所": "seat_copulative_spell",
    "方書": "buillding_name_etc",
    "設置位置": "set_position",
    "緯度": "latitude",
    "経度": "longitude",
    "男性トイレ総数": "male_wc_total",
    "男性トイレ数（小便器）": "male_wc_urinal",
    "男性トイレ数(小便器)": "male_wc_urinal",
    "男性トイレ数（和式）": "male_wc_jstyle",
    "男性トイレ数（洋式）": "male_wc_wstyle",
    "女性トイレ総数": "female_wc_total",
    "女性トイレ数（和式）": "female_wc_jstyle",
    "女性トイレ数（洋式）": "female_wc_wstyle",
    "男女共用トイレ総数": "unisex_wc_total",
    "男女共用トイレ数（和式）": "unisex_wc_jstyle",
    "男女共用トイレ数（洋式）": "unisex_wc_wstyle",
    "車椅子使用者用トイレ有無": "wheelchair_wc",
    "車いす使用者用トイレ有無": "wheelchair_wc",
    "乳幼児用設備設置トイレ有無": "infant_wc",
    "オストメイト設置トイレ有無": "ostomate_wc",
    "利用開始時間": "available_start_time",
    "利用終了時間": "available_end_time",
    "利用可能時間特記事項": "available_time_note",
    "画像": "image",
    "画像_ライセンス": "image_licence",
    "備考": "note",
    "備　　　考": "note",
}

LEGACY = {
    source: PHYSICAL_TO_NAME[target]
    for source, target in LEGACY.items()
}

# 静岡県内の独自形式。
# 原典に明示された意味だけ標準列へ移し、不明・独自項目は備考へ保存する。
CUSTOM = {
    "22225-伊豆の国市": {
        "No": "ID",
        "場所": "名称",
        "住所": "所在地_連結表記",
        "緯度": "緯度",
        "経度": "経度",
        "利用時間": "利用可能時間特記事項",
        "男性トイレ総数": "男性トイレ総数",
        "男性トイレ数（小便）": "男性トイレ数（小便器）",
        "男性トイレ数（和式）": "男性トイレ数（和式）",
        "男性トイレ数（洋式）": "男性トイレ数（洋式）",
        "女性トイレ総数": "女性トイレ総数",
        "女性トイレ数（和式）": "女性トイレ数（和式）",
        "女性トイレ数（洋式）": "女性トイレ数（洋式）",
    },
    "22302-河津町": {
        "No": "ID",
        "名称": "名称",
        "住所": "所在地_連結表記",
        "使用可能時間": "利用可能時間特記事項",
        "緯度": "緯度",
        "経度": "経度",
    },
    "22305-松崎町": {
        "No": "ID",
        "名称": "名称",
        "住所": "所在地_連結表記",
        "使用可能時間": "利用可能時間特記事項",
        "緯度": "緯度",
        "経度": "経度",
    },
}

def s(v):
    return "" if v is None else str(v).strip()

def read_csv(path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    else:
        raise RuntimeError(f"encoding: {path}")

    rows = list(csv.reader(text.splitlines()))

    # 下田などタイトル行付きにも対応。
    hi = next(
        i for i, r in enumerate(rows)
        if "名称" in r or "場所" in r
    )
    header = [s(x) for x in rows[hi]]

    return [
        {header[i]: s(v) for i, v in enumerate(r)
         if i < len(header) and header[i]}
        for r in rows[hi + 1:]
        if any(s(v) for v in r)
    ]

def read_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hi = next(
        i for i, r in enumerate(rows)
        if "名称" in [s(x) for x in r]
    )
    header = [s(x) for x in rows[hi]]

    return [
        {header[i]: s(v) for i, v in enumerate(r)
         if i < len(header) and header[i]}
        for r in rows[hi + 1:]
        if any(s(v) for v in r)
    ]

TARGETS = (
    "22100-静岡市",
    "22203-沼津市",
    "22206-三島市",
    "22207-富士宮市",
    "22208-伊東市",
    "22209-島田市",
    "22211-磐田市",
    "22212-焼津市",
    "22213-掛川市",
    "22214-藤枝市",
    "22215-御殿場市",
    "22216-袋井市",
    "22219-下田市",
    "22220-裾野市",
    "22222-伊豆市",
    "22223-御前崎市",
    "22224-菊川市",
    "22225-伊豆の国市",
    "22226-牧之原市",
    "22301-東伊豆町",
    "22302-河津町",
    "22305-松崎町",
    "22306-西伊豆町",
    "22341-清水町",
)


def main():
    total = 0
    target_filter = set(TARGETS)

    for src in sorted(RAW.glob("*/public-toilet/*")):
        if src.suffix.lower() not in (".csv", ".xlsx"):
            continue

        municipality = src.parents[1].name

        if municipality not in target_filter:
            continue

        rows = read_csv(src) if src.suffix.lower() == ".csv" else read_xlsx(src)

        is_current = rows and "全国地方公共団体コード" in rows[0]
        mapping = CURRENT if is_current else LEGACY

        output = []

        for r in rows:
            # 三島などの ◎/○ メタデータ行。
            # 実施設名ではないことを明示的に判定。
            name = s(r.get("名称"))
            if name in {"◎", "○", "△"}:
                continue

            # IDだけ入った実質空レコードは施設データではないため除外。
            # 「名称が空」という理由だけでは除外しない。
            meaningful = [
                v for k, v in r.items()
                if k not in {
                    "ID", "NO", "No",
                    "全国地方公共団体コード",
                    "都道府県コード又は市区町村コード",
                    "市区町村コード",
                    "地方公共団体名",
                    "都道府県名",
                    "市区町村名",
                }
                and s(v)
            ]
            if not meaningful:
                continue

            out = {field: "" for field in FIELDS}

            custom_mapping = CUSTOM.get(municipality)
            active_mapping = custom_mapping or mapping

            for source_name, target_name in active_mapping.items():
                if target_name in out and source_name in r:
                    out[target_name] = s(r[source_name])

            legacy_code_note = ""

            # 旧データは原典に明示された値だけを標準列へ移す。
            if not is_current and not custom_mapping:
                out["地方公共団体名"] = s(
                    r.get("市区町村名") or r.get("都道府県名")
                )
                out["所在地_都道府県"] = s(r.get("都道府県名"))
                out["所在地_市区町村"] = s(r.get("市区町村名"))

                legacy_code = s(
                    r.get("都道府県コード又は市区町村コード")
                    or r.get("市区町村コード")
                )

                if len(legacy_code) == 6 and legacy_code.isdigit():
                    out["全国地方公共団体コード"] = legacy_code
                elif legacy_code:
                    legacy_code_note = (
                        "都道府県コード又は市区町村コード:" + legacy_code
                    )

            # 多機能トイレ数/有無はバリアフリートイレ数へ推測変換しない。
            extras = []

            if municipality == "22225-伊豆の国市":
                for key in ("多目的トイレ", "手洗い"):
                    if s(r.get(key)):
                        extras.append(f"{key}:{s(r[key])}")

            if municipality in {"22302-河津町", "22305-松崎町"}:
                for key in (
                    "車いす",
                    "オストメイト",
                    "おむつ替えシート",
                    "ベビーチェア",
                ):
                    if s(r.get(key)):
                        extras.append(f"{key}:{s(r[key])}")

            if municipality in {"22302-河津町", "22305-松崎町"}:
                yes_no_mark = {"○": "有", "×": "無"}

                wheelchair = s(r.get("車いす"))
                if wheelchair in yes_no_mark:
                    out["車椅子使用者用トイレ有無"] = yes_no_mark[wheelchair]

                ostomate = s(r.get("オストメイト"))
                if ostomate in yes_no_mark:
                    out["オストメイト設置トイレ有無"] = yes_no_mark[ostomate]

                diaper = s(r.get("おむつ替えシート"))
                if diaper in yes_no_mark:
                    out["乳幼児用設備設置トイレ有無"] = yes_no_mark[diaper]

            if legacy_code_note:
                extras.append(legacy_code_note)

            if s(r.get("多機能トイレ数")):
                extras.append("多機能トイレ数:" + s(r["多機能トイレ数"]))

            if s(r.get("多機能トイレ有無")):
                extras.append("多機能トイレ有無:" + s(r["多機能トイレ有無"]))

            # 伊東固有項目も捨てずnoteへ。
            for key in (
                "点灯時間", "段差", "建築年月",
                "建築面積（㎡）", "建築費（万円）", "処理方法"
            ):
                if s(r.get(key)):
                    extras.append(f"{key}:{s(r[key])}")

            if extras:
                existing = out["備考"]
                out["備考"] = (
                    existing + ("; " if existing else "") + "; ".join(extras)
                )

            # 有無項目は標準語彙「有」「無」のみ採用する。
            # 0/1/2/3等は意味を推測せず、原典値を備考へ保存する。
            for source_field, target_field in (
                ("車椅子使用者用トイレ有無", "車椅子使用者用トイレ有無"),
                ("車いす使用者用トイレ有無", "車椅子使用者用トイレ有無"),
                ("乳幼児用設備設置トイレ有無", "乳幼児用設備設置トイレ有無"),
                ("オストメイト設置トイレ有無", "オストメイト設置トイレ有無"),
            ):
                raw_value = s(r.get(source_field))
                if raw_value and raw_value not in {"有", "無"}:
                    if out.get(target_field) == raw_value:
                        out[target_field] = ""

                    note = out.get("備考", "")
                    preserved = f"{source_field}:{raw_value}"
                    if preserved not in note:
                        out["備考"] = (
                            note + ("; " if note else "") + preserved
                        )

            # 明白な緯度・経度逆転のみ補正する。
            # 日本の緯度・経度の値域から逆転が一意に判断できる場合に限定。
            try:
                lat = float(out.get("緯度", ""))
                lon = float(out.get("経度", ""))
                if 120 <= lat <= 150 and 20 <= lon <= 50:
                    out["緯度"], out["経度"] = out["経度"], out["緯度"]
            except (ValueError, TypeError):
                pass

            # 時刻表記を意味を変えない範囲で HH:MM に統一。
            for time_field in ("利用開始時間", "利用終了時間"):
                value = s(out.get(time_field))

                # 下田の「－」は時刻値ではない。
                if municipality == "22219-下田市" and value == "－":
                    out[time_field] = ""
                    continue

                # 三島の「24時間」は開始時刻ではなく利用可能時間の情報。
                if (
                    municipality == "22206-三島市"
                    and time_field == "利用開始時間"
                    and value == "24時間"
                ):
                    out[time_field] = ""
                    note = s(out.get("利用可能時間特記事項"))
                    if "24時間" not in note:
                        out["利用可能時間特記事項"] = (
                            f"{note}; 24時間" if note else "24時間"
                        )
                    continue

                parts = value.split(":")

                # HH:MM:00 -> HH:MM
                if len(parts) == 3 and parts[2] == "00":
                    parts = parts[:2]

                if len(parts) == 2:
                    try:
                        hour = int(parts[0])
                        minute = int(parts[1])
                        if 0 <= hour <= 24 and 0 <= minute <= 59:
                            out[time_field] = f"{hour:02d}:{minute:02d}"
                    except ValueError:
                        pass

            output.append(out)

        dest = OUT / municipality / "public-toilet/public-toilet.csv"
        dest.parent.mkdir(parents=True, exist_ok=True)

        with dest.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS, lineterminator="\n")
            w.writeheader()
            w.writerows(output)

        print(
            municipality,
            "source=" + str(len(rows)),
            "normalized=" + str(len(output)),
            "type=" + (
                "custom" if custom_mapping
                else ("current39" if is_current else "legacy")
            )
        )

        total += len(output)

    print("total normalized:", total)


if __name__ == "__main__":
    main()
