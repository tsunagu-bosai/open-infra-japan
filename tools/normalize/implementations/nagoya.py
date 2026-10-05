#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from tools import normalize_public_toilet as common
from tools.normalize.core.normalization import (
    AllowedValueStrategy,
    RowSupport,
    StandardCsvWriter,
    StrictHmsTimeStrategy,
)

ROOT = Path(__file__).resolve().parents[3]
SCHEMA = ROOT / "schema/standard/public-toilet/schema.csv"
SOURCE = ROOT / "data/raw/23-愛知県/23100-名古屋市/public-toilet/23100_public-toilet.xlsx"
OUTPUT = ROOT / "data/normalized/23-愛知県/23100-名古屋市/public-toilet/public-toilet.csv"

TARGETS = {"23100": "名古屋市"}

EXPECTED_HEADER = [
    "都道府県コード又は市区町村コード", "NO", "都道府県名", "市区町村名",
    "区名", "名称", "名称_カナ", "名称_英語", "住所", "方書", "設置位置",
    "緯度", "経度", "男性トイレ総数", "男性トイレ数（小便器）",
    "男性トイレ数（和式）", "男性トイレ数（洋式）", "女性トイレ総数",
    "女性トイレ数（和式）", "女性トイレ数（洋式）",
    "男女共用トイレ総数", "男女共用トイレ数（和式）",
    "男女共用トイレ数（洋式）", "多機能トイレ数",
    "車椅子使用者用トイレ有無", "乳幼児用設備設置トイレ有無",
    "ベビーベッド数", "ベビーチェア数",
    "オストメイト設置トイレ有無", "利用開始時間", "利用終了時間",
    "利用可能時間特記事項", "画像", "画像_ライセンス", "備考",
]


def read_schema() -> list[str]:
    with SCHEMA.open("r", encoding="utf-8-sig", newline="") as f:
        return [row["name"] for row in csv.DictReader(f)]


TIME_NORMALIZER = StrictHmsTimeStrategy(cleaner=RowSupport.clean_cell)
BOOLEAN_NORMALIZER = AllowedValueStrategy(("", "有", "無"), cleaner=RowSupport.clean_cell)


def validate_count(value: str, field: str) -> str:
    value = RowSupport.clean_cell(value)
    if value and not value.isdigit():
        raise RuntimeError(f"{field}: non-numeric value {value!r}")
    return value



def main() -> None:
    if "23100" not in TARGETS:
        return
    if OUTPUT.exists():
        raise RuntimeError(f"23100: normalized already exists: {OUTPUT}")

    schema = read_schema()
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    physical = [[RowSupport.clean_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]

    if not physical or physical[0] != EXPECTED_HEADER:
        raise RuntimeError("23100: unexpected source header")

    source_rows = [row for row in physical[1:] if any(RowSupport.clean_cell(v) for v in row)]
    if not source_rows:
        raise RuntimeError("23100: no active source rows")

    normalized = []
    seen_ids: set[str] = set()
    ward_notes = 0
    multi_notes = 0
    baby_detail_notes = 0

    count_fields = [
        "男性トイレ総数", "男性トイレ数（小便器）", "男性トイレ数（和式）",
        "男性トイレ数（洋式）", "女性トイレ総数", "女性トイレ数（和式）",
        "女性トイレ数（洋式）", "男女共用トイレ総数",
        "男女共用トイレ数（和式）", "男女共用トイレ数（洋式）",
    ]

    for values in source_rows:
        src = dict(zip(EXPECTED_HEADER, values))

        if RowSupport.clean_cell(src["都道府県コード又は市区町村コード"]) != "231002":
            raise RuntimeError("23100: unexpected municipality code")

        ident = RowSupport.clean_cell(src["NO"])
        if not ident or ident in seen_ids:
            raise RuntimeError(f"23100: blank/duplicate NO: {ident!r}")
        seen_ids.add(ident)

        row = {field: "" for field in schema}
        row["全国地方公共団体コード"] = "231002"
        row["ID"] = ident
        row["地方公共団体名"] = "愛知県名古屋市"
        row["名称"] = RowSupport.clean_cell(src["名称"])
        row["名称_カナ"] = RowSupport.clean_cell(src["名称_カナ"])
        row["名称_英語"] = RowSupport.clean_cell(src["名称_英語"])
        row["所在地_全国地方公共団体コード"] = "231002"

        ward = RowSupport.clean_cell(src["区名"])
        address = RowSupport.clean_cell(src["住所"])
        prefix = "愛知県名古屋市"
        if address:
            if address.startswith(prefix):
                connected = address
            elif address.startswith("名古屋市"):
                connected = "愛知県" + address
            else:
                connected = prefix + address
        else:
            connected = ""
        row["所在地_連結表記"] = connected
        row["所在地_都道府県"] = "愛知県"
        row["所在地_市区町村"] = "名古屋市"
        row["建物名等(方書)"] = RowSupport.clean_cell(src["方書"])
        row["設置位置"] = RowSupport.clean_cell(src["設置位置"])

        raw_lat = RowSupport.clean_cell(src["緯度"])
        raw_lon = RowSupport.clean_cell(src["経度"])
        if raw_lat or raw_lon:
            raise RuntimeError(
                f"23100: coordinates unexpectedly populated: {raw_lat!r}, {raw_lon!r}"
            )

        for field in count_fields:
            row[field] = validate_count(src[field], field)

        for field in (
            "車椅子使用者用トイレ有無",
            "乳幼児用設備設置トイレ有無",
            "オストメイト設置トイレ有無",
        ):
            row[field] = BOOLEAN_NORMALIZER(src[field], field)

        row["利用開始時間"] = TIME_NORMALIZER(src["利用開始時間"])
        row["利用終了時間"] = TIME_NORMALIZER(src["利用終了時間"])
        row["利用可能時間特記事項"] = RowSupport.clean_cell(src["利用可能時間特記事項"])
        row["画像"] = RowSupport.clean_cell(src["画像"])
        row["画像_ライセンス"] = RowSupport.clean_cell(src["画像_ライセンス"])
        row["備考"] = RowSupport.clean_cell(src["備考"])

        if ward:
            RowSupport.append_note(row, f"原データ区名={ward}")
            ward_notes += 1

        multi = RowSupport.clean_cell(src["多機能トイレ数"])
        if multi:
            validate_count(multi, "多機能トイレ数")
            RowSupport.append_note(row, f"原データ多機能トイレ数={multi}")
            multi_notes += 1

        bed = validate_count(src["ベビーベッド数"], "ベビーベッド数")
        chair = validate_count(src["ベビーチェア数"], "ベビーチェア数")
        if bed not in ("", "0"):
            RowSupport.append_note(row, f"原データベビーベッド数={bed}")
            baby_detail_notes += 1
        if chair not in ("", "0"):
            RowSupport.append_note(row, f"原データベビーチェア数={chair}")
            baby_detail_notes += 1

        normalized.append(row)

    patches = common.load_patches("23-愛知県", "23100-名古屋市")
    applied, problems = common.apply_patches(normalized, patches, SOURCE)
    if problems:
        raise RuntimeError(
            "23100-名古屋市: patch適用失敗: " + " | ".join(map(str, problems))
        )

    StandardCsvWriter().write(OUTPUT, schema, normalized)

    print(
        f"OK 23100 名古屋市: rows={len(normalized)} "
        f"ward_notes={ward_notes} multi_notes={multi_notes} "
        f"baby_detail_notes={baby_detail_notes} patches={applied}"
    )


if __name__ == "__main__":
    main()
