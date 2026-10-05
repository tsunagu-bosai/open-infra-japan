#!/usr/bin/env python3

import csv
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]

SOURCE = (
    ROOT
    / "schema"
    / "standard"
    / "municipal-standard-open-data"
    / "source"
    / "20260801_resources_open_data_municipal-standard-open-dataset_table_a.xlsx"
)

OUTPUT = (
    ROOT
    / "schema"
    / "standard"
    / "public-toilet"
    / "schema.csv"
)

SHEET_NAME = "13.公衆トイレ一覧"

OUTPUT_COLUMNS = [
    "item_no",
    "name",
    "classification",
    "description",
    "format",
    "example",
    "english_name",
    "common_vocabulary",
    "common_vocabulary_type",
    "gif_data_model",
    "gif_classification",
]


def normalize(value):
    if value is None:
        return ""

    return str(value).strip()


def main():
    workbook = load_workbook(
        SOURCE,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook[SHEET_NAME]

    rows = []

    # 4～42行が公衆トイレ一覧の39項目
    for row_number in range(4, 43):
        values = [
            worksheet.cell(
                row=row_number,
                column=column,
            ).value
            for column in range(1, 12)
        ]

        item_no = values[0]

        if item_no is None:
            continue

        rows.append(
            {
                key: normalize(value)
                for key, value in zip(
                    OUTPUT_COLUMNS,
                    values,
                )
            }
        )

    workbook.close()

    if len(rows) != 39:
        raise RuntimeError(
            f"Expected 39 schema rows, got {len(rows)}"
        )

    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with OUTPUT.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=OUTPUT_COLUMNS,
            lineterminator="\n",
        )

        writer.writeheader()
        writer.writerows(rows)

    print(f"Source : {SOURCE}")
    print(f"Sheet  : {SHEET_NAME}")
    print(f"Rows   : {len(rows)}")
    print(f"Output : {OUTPUT}")


if __name__ == "__main__":
    main()
